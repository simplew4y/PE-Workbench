"""Preserve the main workbook regressions through the pure parser transport.

SQLite cache publication/rollback is covered by test/excel-processing.test.ts,
because the Node service now owns those transactions.
"""
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName

import workbook_reader


class DocumentCacheTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='pe-reader-cache-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / 'model_2025_Jun_30.xlsx'
        workbook = Workbook()
        workbook.epoch = CALENDAR_MAC_1904
        sheet = workbook.active
        sheet.title = 'Model'
        sheet.append(['Valuation Date', date(2025, 6, 30)])
        sheet.append(['Forecast', '2030E'])
        sheet.append(['Target price date', date(2026, 6, 30)])
        sheet.append(['Valuation Date', '=DATE(2025,6,30)'])
        workbook.defined_names.add(DefinedName('Valuation_Date', attr_text="'Model'!$B$1"))
        workbook.save(self.source)
        workbook.close()

    def parse(self, revision='r1'):
        output = self.root / (revision + '.json')
        output.unlink(missing_ok=True)
        subprocess.run([sys.executable, str(Path(__file__).with_name('parse_workbook.py')),
                        '--input', str(self.source), '--output', str(output),
                        '--doc-id', 'a' * 40, '--dataset-id', 'dataset-1', '--revision', revision,
                        '--sha256', hashlib.sha256(self.source.read_bytes()).hexdigest(),
                        '--filename', self.source.name], check=True)
        return json.loads(output.read_text())

    def test_keeps_1904_dates_named_ranges_and_forecasts_separate(self):
        result = self.parse()
        # Preparation stores navigation, not inferred financial dates or all cells.
        self.assertEqual(result['tables']['valuation_date_candidates'], [])
        self.assertEqual(result['tables']['excel_cells'], [])
        self.assertEqual(result['document_date'], '')
        cells = workbook_reader.read_workbook(self.source, {
            'action': 'read', 'sheet': 'Model', 'range': 'B1:B4',
        }, navigation=result['navigation'])['cells']
        self.assertEqual([cell['value'] for cell in cells], [
            '2025-06-30T00:00:00', '2030E', '2026-06-30T00:00:00', '=DATE(2025,6,30)',
        ])
        self.assertEqual(cells[3]['formula_cache_status'], 'missing')
        name = result['navigation']['defined_names'][0]
        self.assertEqual(name['name'], 'Valuation_Date')
        self.assertEqual(name['destinations'], [['Model', '$B$1']])

    def test_refresh_preserves_navigation_and_original_even_after_a_parser_failure(self):
        first = self.parse()
        checksum = hashlib.sha256(self.source.read_bytes()).hexdigest()
        with patch.object(workbook_reader, 'inspect_workbook', side_effect=RuntimeError('refresh failed')):
            with self.assertRaisesRegex(RuntimeError, 'refresh failed'):
                workbook_reader.navigation_artifact(self.source, 'dataset-1', 'a' * 40)
        second = self.parse('r2')
        self.assertEqual(first['tables'], second['tables'])
        self.assertEqual(first['navigation'], second['navigation'])
        self.assertEqual(first['doc_id'], second['doc_id'])
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), checksum)

    def test_defined_name_preserves_qualification_on_its_target(self):
        workbook = load_workbook(self.source)
        workbook['Model']['A1'] = 'Valuation Date not confirmed'
        workbook.save(self.source)
        workbook.close()
        result = self.parse()
        self.assertEqual(result['tables']['valuation_date_candidates'], [])
        self.assertEqual(result['navigation']['defined_names'][0]['destinations'], [['Model', '$B$1']])
        cells = workbook_reader.read_workbook(self.source, {
            'action': 'read', 'sheet': 'Model', 'range': 'A1:B1',
        }, navigation=result['navigation'])['cells']
        self.assertEqual(cells[0]['value'], 'Valuation Date not confirmed')
        self.assertEqual(cells[1]['value'], '2025-06-30T00:00:00')
        self.assertNotIn('assertion_status', cells[1])


if __name__ == '__main__':
    unittest.main()
