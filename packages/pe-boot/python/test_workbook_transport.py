"""Exercise the standalone parser without a project database or runtime caches."""
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from validate_workbook import validate_workbook
from workbook import parse_workbook


def workbook_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Valuation"
    sheet.append(["Target Price", "=B2*2"])
    sheet.append(["Current Price", 10])
    sheet["B1"].number_format = '"CNY/share" 0.00'
    sheet.merge_cells("C4:D4")
    sheet["C4"] = "Merged cell"
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "=Valuation!B2"
    hidden.sheet_state = "veryHidden"
    workbook.defined_names.add(DefinedName("Primary_Target", attr_text="'Valuation'!$B$1"))
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class WorkbookTransportTests(unittest.TestCase):
    def test_rejects_non_workbooks_and_wrong_ooxml_type(self):
        with self.assertRaises(ValueError):
            validate_workbook(b"PK renamed arbitrary data", "xlsx")
        with self.assertRaisesRegex(ValueError, "content type"):
            validate_workbook(workbook_bytes(), "xlsm")

    def test_accepts_xlsx_and_macro_enabled_container_without_executing_macros(self):
        data = workbook_bytes()
        validate_workbook(data, "xlsx")
        with zipfile.ZipFile(io.BytesIO(data)) as original:
            entries = {name: original.read(name) for name in original.namelist()}
        entries["[Content_Types].xml"] = entries["[Content_Types].xml"].replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        )
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        validate_workbook(output.getvalue(), "xlsm")

    def test_rejects_zip_with_missing_workbook_relationship(self):
        with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as original:
            entries = {name: original.read(name) for name in original.namelist() if name != "_rels/.rels"}
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        with self.assertRaisesRegex(ValueError, "missing"):
            validate_workbook(output.getvalue(), "xlsx")

    def test_pure_parser_preserves_formulas_hidden_state_names_and_json_scalar_types(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Model.xlsx"
            path.write_bytes(workbook_bytes())
            result = parse_workbook(
                dataset_id="dataset", doc_id="a" * 40, path=path,
                source_modified_at="2026-09-07T00:00:00.123Z",
            )
            self.assertEqual(set(result["tables"]), {
                "excel_workbooks", "excel_sheets", "excel_regions", "excel_cells",
                "excel_defined_names", "excel_formula_references", "valuation_date_candidates", "metric_facts",
            })
            cells = result["tables"]["excel_cells"]
            formula = next(cell for cell in cells if cell["sheet_name"] == "Valuation" and cell["cell_ref"] == "B1")
            self.assertEqual(formula["formula"], "=B2*2")
            self.assertIsNone(formula["cached_value"])
            self.assertTrue(any(sheet["sheet_state"] == "veryHidden" for sheet in result["tables"]["excel_sheets"]))
            self.assertTrue(result["tables"]["excel_defined_names"])
            self.assertTrue(result["tables"]["excel_formula_references"])
            json.dumps(result, allow_nan=False)
            self.assertEqual([item.name for item in Path(directory).iterdir()], ["Model.xlsx"])


if __name__ == "__main__":
    unittest.main()
