#!/usr/bin/env python3
"""Apply approved literal-cell changes and inspect a recalculated workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.dom import minidom

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, coordinate_to_tuple, range_boundaries

from excel_formula_parser import extract_formula_references


CELL_RE = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
EXTERNAL_FORMULA_RE = re.compile(r"_xll\.|_xludf\.|\[(?:\d+|[^\]]+\.xls[xmb]?)\]", re.IGNORECASE)


def validate_preserved_caches(workbook: Any, changes: list[dict[str, Any]]) -> None:
    """Only keep unavailable add-in results whose inputs were not changed."""
    changed = [(item["sheet"], coordinate_to_tuple(item["cell"])) for item in changes if not same_value(item["before"], item["after"])]
    if not changed:
        return
    formulas = {sheet.title: {cell.coordinate: cell.value for cell in sheet._cells.values() if cell.data_type == "f"} for sheet in workbook.worksheets}
    visited: set[tuple[str, str]] = set()

    def inspect(sheet_name: str, formula: str) -> None:
        key = (sheet_name, formula)
        if key in visited:
            return
        visited.add(key)
        if re.search(r"\b(?:INDIRECT|OFFSET)\s*\(", formula, re.IGNORECASE):
            raise ValueError("Cannot verify dynamic dependencies of an unavailable formula cache")
        names = dict(workbook.defined_names)
        names.update(workbook[sheet_name].defined_names)
        for ref in extract_formula_references(formula, source_sheet=sheet_name, defined_names=names):
            if ref.external_workbook:
                continue
            if ref.reference_kind == "defined_name":
                defined = next((value for name, value in names.items() if name.casefold() == str(ref.defined_name).casefold()), None)
                if defined is None:
                    raise ValueError("Cannot resolve an unavailable formula cache dependency")
                inspect(sheet_name, "=" + str(defined.attr_text).lstrip("="))
                continue
            if ref.parse_status != "resolved" or not ref.target_range or ref.target_sheet not in formulas:
                raise ValueError(f"Cannot verify unavailable formula cache dependency: {ref.raw_reference}")
            left, top, right, bottom = range_boundaries(ref.target_range)
            def contains(position: tuple[int, int]) -> bool:
                row, column = position
                return (left is None or column >= left) and (right is None or column <= right) and (top is None or row >= top) and (bottom is None or row <= bottom)
            if any(sheet == ref.target_sheet and contains(position) for sheet, position in changed):
                raise ValueError(f"Changed input affects an unavailable formula cache: {ref.target_sheet}!{ref.target_range}")
            for address, dependency in formulas[ref.target_sheet].items():
                if contains(coordinate_to_tuple(address)):
                    if not isinstance(dependency, str):
                        raise ValueError("Cannot verify array formula dependency of an unavailable cache")
                    inspect(ref.target_sheet, dependency)

    for sheet_name, cells in formulas.items():
        for address, formula in cells.items():
            expression = formula if isinstance(formula, str) else getattr(formula, "text", None)
            if isinstance(expression, str) and EXTERNAL_FORMULA_RE.search(expression):
                try:
                    inspect(sheet_name, expression)
                except ValueError as error:
                    raise ValueError(f"Cannot preserve {sheet_name}!{address}: {error}") from error


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError("Model update plan must be a JSON object")
    return value


def write_json(path: Path, value: dict[str, Any]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def workbook_options(path: Path, *, data_only: bool, read_only: bool) -> dict[str, Any]:
    return {
        "data_only": data_only,
        "keep_links": True,
        "keep_vba": path.suffix.lower() == ".xlsm",
        "read_only": read_only,
    }


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def same_value(left: Any, right: Any) -> bool:
    if isinstance(left, (int, float)) and not isinstance(left, bool) and isinstance(right, (int, float)) and not isinstance(right, bool):
        return math.isclose(float(left), float(right), rel_tol=1e-12, abs_tol=1e-12)
    return left == right


def validate_cell_ref(cell_ref: str) -> str:
    normalized = cell_ref.strip().upper().replace("$", "")
    if not CELL_RE.fullmatch(normalized):
        raise ValueError(f"Invalid A1 cell reference: {cell_ref}")
    coordinate_from_string(normalized)
    return normalized


def formula_token(value: Any) -> str:
    if isinstance(value, str):
        return value
    attributes = getattr(value, "__dict__", None)
    if isinstance(attributes, dict):
        return f"{type(value).__name__}:{json.dumps(attributes, sort_keys=True, separators=(',', ':'))}"
    return f"{type(value).__name__}:{value!s}"


def formula_fingerprint(workbook: Any) -> tuple[str, int]:
    digest = hashlib.sha256()
    count = 0
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                count += 1
                digest.update(
                    f"{worksheet.title}\0{cell.coordinate}\0{formula_token(cell.value)}\n".encode("utf-8")
                )
    return digest.hexdigest(), count


def formula_map(workbook: Any) -> dict[str, str]:
    formulas: dict[str, str] = {}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas[f"{worksheet.title}!{cell.coordinate}"] = formula_token(cell.value)
    return formulas


def error_cells(workbook: Any) -> list[str]:
    errors: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and value.startswith("#")):
                    errors.append(f"{worksheet.title}!{cell.coordinate}:{value}")
    return sorted(errors)


def apply_changes(source: Path, plan_path: Path, destination: Path, result_path: Path) -> None:
    if source.suffix.lower() not in {".xlsx", ".xlsm"} or destination.suffix.lower() != source.suffix.lower():
        raise ValueError("Controlled model updates preserve the source .xlsx or .xlsm format")
    plan = load_json(plan_path)
    changes = plan.get("changes")
    if not isinstance(changes, list) or not changes:
        raise ValueError("Model update plan must contain changes")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, **workbook_options(source, data_only=False, read_only=False))
    try:
        before_fingerprint, formula_count = formula_fingerprint(workbook)
        seen: set[tuple[str, str]] = set()
        applied: list[dict[str, Any]] = []
        for change in changes:
            if not isinstance(change, dict):
                raise ValueError("Model update change must be an object")
            sheet_name = str(change.get("sheet", ""))
            cell_ref = validate_cell_ref(str(change.get("cell", "")))
            key = (sheet_name, cell_ref)
            if key in seen:
                raise ValueError(f"Duplicate model update cell: {sheet_name}!{cell_ref}")
            seen.add(key)
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            cell = workbook[sheet_name][cell_ref]
            if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                raise ValueError(f"Refusing to overwrite formula cell {sheet_name}!{cell_ref}")
            expected = change.get("expected_value")
            before = json_value(cell.value)
            if not same_value(before, expected):
                raise ValueError(
                    f"Expected value mismatch at {sheet_name}!{cell_ref}: expected {expected!r}, found {before!r}"
                )
            new_value = change.get("new_value")
            if not (new_value is None or isinstance(new_value, (str, int, float, bool))):
                raise ValueError(f"Unsupported value at {sheet_name}!{cell_ref}")
            cell.value = new_value
            applied.append(
                {
                    "sheet": sheet_name,
                    "cell": cell_ref,
                    "before": before,
                    "after": new_value,
                    "kind": str(change.get("kind", "")),
                    "metric_id": change.get("metric_id"),
                    "evidence_ids": change.get("evidence_ids", []),
                    "reason": str(change.get("reason", "")),
                }
            )
        after_fingerprint, after_formula_count = formula_fingerprint(workbook)
        if before_fingerprint != after_fingerprint or formula_count != after_formula_count:
            raise ValueError("Formula cells changed while applying literal input updates")
        validate_preserved_caches(workbook, applied)
        # Saving the entire book through openpyxl clears every formula cache,
        # including Bloomberg add-in values unavailable to LibreOffice. Patch
        # approved literal cells in the original ZIP, and invalidate ordinary
        # formula caches so the calculation engine refreshes their results.
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, "w") as target:
            book_xml = minidom.parseString(original.read("xl/workbook.xml"))
            rels_xml = minidom.parseString(original.read("xl/_rels/workbook.xml.rels"))
            relations = {node.getAttribute("Id"): node.getAttribute("Target") for node in rels_xml.getElementsByTagNameNS("*", "Relationship")}
            sheet_paths = {}
            for node in book_xml.getElementsByTagNameNS("*", "sheet"):
                path = relations[node.getAttributeNS("http://schemas.openxmlformats.org/officeDocument/2006/relationships", "id")]
                sheet_paths[node.getAttribute("name")] = path.lstrip("/") if path.startswith("/") else posixpath.normpath(posixpath.join("xl", path))
            book_xml.unlink()
            rels_xml.unlink()
            patches = {}
            for change in applied:
                patches.setdefault(sheet_paths[change["sheet"]], []).append(change)
            ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            for entry in original.infolist():
                data = original.read(entry.filename)
                if entry.filename in sheet_paths.values():
                    document = minidom.parseString(data)
                    def element(name: str) -> Any:
                        prefix = document.documentElement.prefix
                        return document.createElementNS(ns, f"{prefix}:{name}" if prefix else name)
                    sheet_data_nodes = document.getElementsByTagNameNS(ns, "sheetData")
                    if not sheet_data_nodes:
                        document.unlink()
                        target.writestr(entry, data)
                        continue
                    # Recompute ordinary formulas, but preserve add-in caches that
                    # cannot be fetched by the local calculation engine.
                    for formula in document.getElementsByTagNameNS(ns, "f"):
                        expression = "".join(node.data for node in formula.childNodes if node.nodeType in {node.TEXT_NODE, node.CDATA_SECTION_NODE})
                        if EXTERNAL_FORMULA_RE.search(expression):
                            continue
                        for node in list(formula.parentNode.childNodes):
                            if node.nodeType == node.ELEMENT_NODE and node.localName == "v":
                                formula.parentNode.removeChild(node)
                    sheet_data = sheet_data_nodes[0]
                    for change in patches.get(entry.filename, []):
                        ref = change["cell"]
                        row_number = int(re.search(r"\d+$", ref).group())
                        row = next((node for node in sheet_data.childNodes if node.nodeType == node.ELEMENT_NODE and node.getAttribute("r") == str(row_number)), None)
                        if row is None:
                            row = element("row")
                            row.setAttribute("r", str(row_number))
                            before = next((node for node in sheet_data.childNodes if node.nodeType == node.ELEMENT_NODE and int(node.getAttribute("r")) > row_number), None)
                            sheet_data.insertBefore(row, before)
                        cell = next((node for node in row.childNodes if node.nodeType == node.ELEMENT_NODE and node.getAttribute("r") == ref), None)
                        if cell is None:
                            cell = element("c")
                            cell.setAttribute("r", ref)
                            column = coordinate_from_string(ref)[0]
                            before = next((node for node in row.childNodes if node.nodeType == node.ELEMENT_NODE and (len(coordinate_from_string(node.getAttribute("r"))[0]), coordinate_from_string(node.getAttribute("r"))[0]) > (len(column), column)), None)
                            row.insertBefore(cell, before)
                        for node in list(cell.childNodes):
                            if node.nodeType == node.ELEMENT_NODE and node.localName in {"v", "is"}:
                                cell.removeChild(node)
                        value = change["after"]
                        if cell.hasAttribute("t"):
                            cell.removeAttribute("t")
                        if value is not None:
                            if isinstance(value, str):
                                cell.setAttribute("t", "inlineStr")
                                inline = element("is")
                                text = element("t")
                                text.setAttribute("xml:space", "preserve")
                                text.appendChild(document.createTextNode(value))
                                inline.appendChild(text)
                                before = next((node for node in cell.childNodes if node.nodeType == node.ELEMENT_NODE and node.localName == "extLst"), None)
                                cell.insertBefore(inline, before)
                            else:
                                if isinstance(value, bool):
                                    cell.setAttribute("t", "b")
                                node = element("v")
                                node.appendChild(document.createTextNode(str(int(value)) if isinstance(value, bool) else str(value)))
                                before = next((child for child in cell.childNodes if child.nodeType == child.ELEMENT_NODE and child.localName == "extLst"), None)
                                cell.insertBefore(node, before)
                    dimensions = document.getElementsByTagNameNS(ns, "dimension")
                    if dimensions and entry.filename in patches:
                        dimensions[0].setAttribute("ref", workbook[patches[entry.filename][0]["sheet"]].calculate_dimension())
                    data = document.toxml(encoding="utf-8")
                    document.unlink()
                target.writestr(entry, data)
        write_json(
            result_path,
            {
                "applied_changes": applied,
                "formula_fingerprint": before_fingerprint,
                "formula_cell_count": formula_count,
                "source_error_cells": error_cells(workbook),
                "sheet_names": workbook.sheetnames,
            },
        )
    finally:
        workbook.close()


def inspect_workbooks(source: Path, updated: Path, original: Path, plan_path: Path, result_path: Path) -> None:
    if original.suffix.lower() == ".xlsm":
        with zipfile.ZipFile(original) as original_archive, zipfile.ZipFile(updated) as updated_archive:
            for name in original_archive.namelist():
                if "vba" in name.lower() and (name not in updated_archive.namelist() or original_archive.read(name) != updated_archive.read(name)):
                    raise ValueError(f"Workbook VBA content changed during recalculation: {name}")
    plan = load_json(plan_path)
    outputs = plan.get("outputs", [])
    checks = plan.get("checks", [])
    if not isinstance(outputs, list) or not isinstance(checks, list):
        raise ValueError("Outputs and checks must be arrays")
    source_formulas = load_workbook(source, **workbook_options(source, data_only=False, read_only=True))
    source_values = load_workbook(source, **workbook_options(source, data_only=True, read_only=True))
    updated_formulas = load_workbook(updated, **workbook_options(updated, data_only=False, read_only=True))
    updated_values = load_workbook(updated, **workbook_options(updated, data_only=True, read_only=True))
    try:
        source_errors = error_cells(source_values)
        updated_errors = error_cells(updated_values)
        observed_outputs: dict[str, Any] = {}
        for output in outputs:
            if not isinstance(output, dict):
                raise ValueError("Output observation must be an object")
            sheet_name = str(output.get("sheet", ""))
            cell_ref = validate_cell_ref(str(output.get("cell", "")))
            key = str(output.get("key", ""))
            if not key:
                raise ValueError("Output observation requires a key")
            if sheet_name not in updated_values.sheetnames:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            observed_outputs[key] = json_value(updated_values[sheet_name][cell_ref].value)
        check_results: list[dict[str, Any]] = []
        for check in checks:
            if not isinstance(check, dict):
                raise ValueError("Validation check must be an object")
            sheet_name = str(check.get("sheet", ""))
            cell_ref = validate_cell_ref(str(check.get("cell", "")))
            key = str(check.get("key", ""))
            expected = check.get("expected_value")
            tolerance = float(check.get("tolerance", 0))
            actual = json_value(updated_values[sheet_name][cell_ref].value)
            if isinstance(actual, (int, float)) and not isinstance(actual, bool) and isinstance(expected, (int, float)) and not isinstance(expected, bool):
                passed = abs(float(actual) - float(expected)) <= tolerance
            else:
                passed = actual == expected
            check_results.append(
                {
                    "key": key,
                    "sheet": sheet_name,
                    "cell": cell_ref,
                    "actual": actual,
                    "expected": expected,
                    "tolerance": tolerance,
                    "passed": passed,
                }
            )
        source_fingerprint, source_formula_count = formula_fingerprint(source_formulas)
        fingerprint, formula_count = formula_fingerprint(updated_formulas)
        source_formula_map = formula_map(source_formulas)
        updated_formula_map = formula_map(updated_formulas)
        formula_changes = sorted(
            cell_ref
            for cell_ref in set(source_formula_map) | set(updated_formula_map)
            if source_formula_map.get(cell_ref) != updated_formula_map.get(cell_ref)
        )
        write_json(
            result_path,
            {
                "outputs": observed_outputs,
                "checks": check_results,
                "source_error_cells": source_errors,
                "updated_error_cells": updated_errors,
                "new_error_cells": sorted(set(updated_errors) - set(source_errors)),
                "source_formula_fingerprint": source_fingerprint,
                "source_formula_cell_count": source_formula_count,
                "formula_fingerprint": fingerprint,
                "formula_cell_count": formula_count,
                "formula_changed_cells": formula_changes,
                "sheet_names": updated_formulas.sheetnames,
            },
        )
    finally:
        source_formulas.close()
        source_values.close()
        updated_formulas.close()
        updated_values.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    apply_parser = subparsers.add_parser("apply")
    apply_parser.add_argument("--input", required=True, type=Path)
    apply_parser.add_argument("--plan", required=True, type=Path)
    apply_parser.add_argument("--output", required=True, type=Path)
    apply_parser.add_argument("--result", required=True, type=Path)

    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--source", required=True, type=Path)
    inspect_parser.add_argument("--updated", required=True, type=Path)
    inspect_parser.add_argument("--original", required=True, type=Path)
    inspect_parser.add_argument("--plan", required=True, type=Path)
    inspect_parser.add_argument("--result", required=True, type=Path)

    args = parser.parse_args()
    if args.command == "apply":
        apply_changes(args.input, args.plan, args.output, args.result)
    else:
        inspect_workbooks(args.source, args.updated, args.original, args.plan, args.result)


if __name__ == "__main__":
    main()
