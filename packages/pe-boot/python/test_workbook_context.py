from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest
from xml.etree import ElementTree
import zipfile

from openpyxl import Workbook

from workbook import parse_workbook


class WorkbookContextTest(unittest.TestCase):
    def parse(self, workbook: Workbook, caches: dict[str, str] | None = None) -> dict:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            workbook.save(path)
            workbook.close()
            if caches:
                archive_path = Path(directory) / "cached.xlsx"
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                with zipfile.ZipFile(path) as source, zipfile.ZipFile(archive_path, "w") as target:
                    for entry in source.infolist():
                        data = source.read(entry.filename)
                        if entry.filename == "xl/worksheets/sheet1.xml":
                            root = ElementTree.fromstring(data)
                            for cell in root.iter(namespace + "c"):
                                if cell.attrib["r"] in caches:
                                    cell.set("t", "str")
                                    value = cell.find(namespace + "v")
                                    if value is None:
                                        value = ElementTree.SubElement(cell, namespace + "v")
                                    value.text = caches[cell.attrib["r"]]
                            data = ElementTree.tostring(root)
                        target.writestr(entry, data)
                path = archive_path
            return parse_workbook(dataset_id="synthetic", doc_id="doc_" + "b" * 24, path=path)

    def cell(self, result: dict, ref: str) -> dict:
        return next(row for row in result["tables"]["excel_cells"] if row["cell_ref"] == ref)

    def test_history_forecast_and_units_survive_column_and_row_shifts(self) -> None:
        for row_shift, col_shift in ((0, 1), (7, 13)):
            with self.subTest(row_shift=row_shift, col_shift=col_shift):
                workbook = Workbook()
                sheet = workbook.active
                sheet.cell(1 + row_shift, 1, "Growth %")
                sheet.cell(2 + row_shift, 1, "Example company")
                sheet.cell(3 + row_shift, 1, "Year end December; EUR m")
                caches = {}
                for index, year in enumerate((1993, 1994, 1995, 2023, 2024, 2025, 2026)):
                    col = 1 + col_shift + index
                    sheet.cell(2 + row_shift, col, year)
                    header = sheet.cell(3 + row_shift, col, f'={year}&"E"' if year >= 2025 else year)
                    if year >= 2025:
                        caches[header.coordinate] = f"{year}E"
                        sheet.cell(4 + row_shift, col, 0)
                    sheet.cell(6 + row_shift, col, 500 + index * 100)
                    sheet.cell(7 + row_shift, col, 0.65).number_format = "0.0%"
                    sheet.cell(8 + row_shift, col, 105)
                sheet.cell(6 + row_shift, 1, "Group sales")
                sheet.cell(7 + row_shift, 1, "Gross margin")
                sheet.cell(8 + row_shift, 1, "Weighted average diluted shares (m)")
                result = self.parse(workbook, caches)
                historical = self.cell(result, sheet.cell(6 + row_shift, 3 + col_shift).coordinate)
                current = self.cell(result, sheet.cell(6 + row_shift, 5 + col_shift).coordinate)
                forecast = self.cell(result, sheet.cell(6 + row_shift, 7 + col_shift).coordinate)
                self.assertEqual((historical["period"], historical["unit"]), ("1995", "EURm"))
                self.assertEqual((current["period"], current["unit"]), ("2024", "EURm"))
                self.assertEqual((forecast["period"], forecast["col_label"]), ("2026E", "2026E"))
                self.assertEqual(self.cell(result, sheet.cell(7 + row_shift, 5 + col_shift).coordinate)["unit"], "%")
                self.assertEqual(self.cell(result, sheet.cell(8 + row_shift, 5 + col_shift).coordinate)["unit"], "shares_m")
                metadata = json.loads(current["metadata_json"])
                self.assertEqual(metadata["period_context"]["sources"][0]["cell_ref"], sheet.cell(3 + row_shift, 5 + col_shift).coordinate)
                self.assertEqual(metadata["unit_context"]["sources"][0]["cell_ref"], f"A{3 + row_shift}")
                self.assertEqual(metadata["verification_status"], "unverified")

    def test_valuation_summary_does_not_inherit_forecast_year(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "€m"
        for col, year in enumerate((2023, 2024, 2025), start=2):
            sheet.cell(3, col, year)
            sheet.cell(4, col, 1000 + col)
        sheet["A4"] = "Free cash flow"
        sheet["A7"] = "Valuation summary"
        sheet["A8"] = "Equity value per Share (€)"
        sheet["C8"] = 2200
        result = self.parse(workbook)
        output = self.cell(result, "C8")
        self.assertEqual(output["period"], "")
        self.assertEqual(output["col_label"], "")
        self.assertEqual(output["unit"], "EUR/share")
        context = json.loads(output["metadata_json"])["period_context"]
        self.assertEqual(context["status"], "missing")
        self.assertEqual(context["reason"], "outside_header_data_layout")
        fact = next(row for row in result["tables"]["metric_facts"] if row["cell_ref"] == "C8")
        self.assertEqual(fact["quality_status"], "review_required")

    def test_two_year_table_does_not_label_single_value_output(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["EURm", 2024, 2025])
        sheet.append(["Revenue", 100, 200])
        sheet.append(["Current share price", None, 2024])
        result = self.parse(workbook)
        output = self.cell(result, "C3")
        self.assertEqual((output["period"], output["unit"]), ("", "EUR/share"))

    def test_numeric_amount_equal_to_year_does_not_replace_header(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Year", 2023])
        sheet.append(["Revenue EURm", 2024])
        sheet.append(["Net profit EURm", 50])
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B2")["period"], "2023")
        self.assertEqual(self.cell(result, "B3")["period"], "2023")
        self.assertTrue(any(row["cell_ref"] == "B2" for row in result["tables"]["metric_facts"]))

    def test_year_shaped_amounts_in_an_existing_table_remain_data(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["EURm", 2023, 2024])
        sheet.append(["Revenue", 2026, 2027])
        sheet.append(["Net profit", 50, 60])
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B2")["period"], "2023")
        self.assertEqual(self.cell(result, "C3")["period"], "2024")
        self.assertTrue(any(row["cell_ref"] == "B2" for row in result["tables"]["metric_facts"]))

    def test_independent_input_panel_does_not_invalidate_annual_headers(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["EURm", 2024, 2025, None, "Assumed multiple", 47])
        sheet.append(["Revenue", 100, 200, None, "Current share price", 250])
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B2")["period"], "2024")
        self.assertEqual(self.cell(result, "C2")["period"], "2025")
        self.assertEqual(self.cell(result, "F2")["period"], "")

    def test_separate_tables_preserve_local_currency_and_new_periods(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        for ref, value in {"A1": "EURm", "H1": "USDm", "B2": 2023, "C2": 2024,
                           "I2": 2025, "J2": 2026, "A3": "Sales", "B3": 10, "C3": 11,
                           "H3": "Sales", "I3": 20, "J3": 21, "A5": "EURm",
                           "A8": "Sales", "B8": 12, "C8": 13, "H8": "Sales", "I8": 22, "J8": 23,
                           "B10": "GBPm", "C11": 2029, "D11": 2030,
                           "B12": "Revenue", "C12": 30, "D12": 31}.items():
            sheet[ref] = value
        result = self.parse(workbook)
        self.assertEqual((self.cell(result, "C8")["period"], self.cell(result, "C8")["unit"]), ("2024", "EURm"))
        self.assertEqual((self.cell(result, "J8")["period"], self.cell(result, "J8")["unit"]), ("2026", "USDm"))
        self.assertEqual((self.cell(result, "C12")["period"], self.cell(result, "C12")["unit"]), ("2029", "GBPm"))

    def test_percent_label_never_becomes_sheet_unit(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Growth %", 0.1])
        sheet.append(["Year", 2024, 2025])
        sheet.append(["Revenue", 100, 200])
        sheet.append(["Margin", 0.2, 0.3])
        sheet["B4"].number_format = "0.0%"
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B3")["unit"], "")
        self.assertEqual(self.cell(result, "B4")["unit"], "%")

    def test_units_preserve_currency_scale_and_denominator(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        expected = [
            ("Revenue (Eur m)", "EURm"), ("Cash (€m)", "EURm"),
            ("Cash EUR", "EUR"), ("EPS EUR/share", "EUR/share"),
            ("Earnings per share (€)", "EUR/share"), ("Total shares (m)", "shares_m"),
            ("Shares outstanding (shares)", "shares"), ("Assumed multiple", "x"),
            ("P/E", "x"), ("EV/EBITDA (x)", "x"),
        ]
        for row, (label, _) in enumerate(expected, start=1):
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, 12)
        result = self.parse(workbook)
        for row, (_, unit) in enumerate(expected, start=1):
            self.assertEqual(self.cell(result, f"B{row}")["unit"], unit)

    def test_merged_year_header_retains_original_source_coordinate(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet["B1"] = 2024
        sheet.merge_cells("B1:D1")
        sheet.append(["Sales EURm", 10, 20, 30])
        result = self.parse(workbook)
        cell = self.cell(result, "C2")
        self.assertEqual(cell["period"], "2024")
        self.assertEqual(json.loads(cell["metadata_json"])["period_context"]["sources"][0]["cell_ref"], "B1")

    def test_conflicting_local_units_require_review(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Year", 2024])
        sheet.append(["Revenue EURm", 10])
        sheet["B2"].number_format = "0.0%"
        result = self.parse(workbook)
        cell = self.cell(result, "B2")
        self.assertEqual(cell["unit"], "")
        self.assertEqual(json.loads(cell["metadata_json"])["unit_context"]["status"], "ambiguous")
        fact = next(row for row in result["tables"]["metric_facts"] if row["cell_ref"] == "B2")
        self.assertEqual(fact["quality_status"], "review_required")

    def test_price_currency_from_number_format_is_preserved(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Target price", 120])
        sheet["B1"].number_format = '"CNY/share" 0.00'
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B1")["unit"], "CNY/share")

    def test_share_count_scale_requires_explicit_evidence_not_value_size(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Year", 2024, 2025])
        sheet.append(["Number of shares (million)", 100, 101])
        sheet.append(["Number of shares (individual shares)", 100, 101])
        sheet.append(["Number of shares", 100, 100_000_000])
        sheet.append(["Shares", 100, 101])
        sheet["B5"].number_format = '0 "shares"'
        result = self.parse(workbook)
        self.assertEqual(self.cell(result, "B2")["unit"], "shares_m")
        self.assertEqual(self.cell(result, "B3")["unit"], "shares")
        for ref in ("B4", "C4", "C5"):
            self.assertEqual(self.cell(result, ref)["unit"], "share_count_unspecified_scale")
            self.assertEqual(json.loads(self.cell(result, ref)["metadata_json"])["unit_context"]["status"], "missing")
        self.assertEqual(self.cell(result, "B5")["unit"], "shares")
        unknown = next(row for row in result["tables"]["metric_facts"] if row["cell_ref"] == "B4")
        self.assertEqual(unknown["quality_status"], "review_required")

    def test_complete_candidate_is_explicitly_unverified(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Year", 2024])
        sheet.append(["Revenue EURm", 10])
        result = self.parse(workbook)
        fact = next(row for row in result["tables"]["metric_facts"] if row["cell_ref"] == "B2")
        self.assertEqual(fact["quality_status"], "candidate_complete")
        self.assertEqual(fact["fact_status"], "candidate")
        metadata = json.loads(fact["metadata_json"])
        self.assertEqual(metadata["verification_status"], "unverified")
        self.assertEqual(metadata["quality_status_scope"], "extraction_completeness_only")


if __name__ == "__main__":
    unittest.main()
