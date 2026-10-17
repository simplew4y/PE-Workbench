import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.formula import ArrayFormula

from model_update_workbook import apply_changes


class ModelUpdateTests(unittest.TestCase):
    def test_rejects_changed_external_cache_dependencies(self):
        for formula in ("=_xll.UNKNOWN(A1)", "=_xll.UNKNOWN(C1)", ArrayFormula(ref="B1", text="=_xll.UNKNOWN(A1)")):
            with self.subTest(formula=formula), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                book = Workbook()
                book.active["A1"] = 100
                book.active["C1"] = "=A1*2"
                book.active["B1"] = formula
                book.save(root / "source.xlsx")
                (root / "plan.json").write_text(json.dumps({"changes": [{"sheet": "Sheet", "cell": "A1", "expected_value": 100, "new_value": 120}]}))
                with self.assertRaisesRegex(ValueError, "Changed input affects"):
                    apply_changes(root / "source.xlsx", root / "plan.json", root / "result.xlsx", root / "result.json")

    def test_preserves_unrelated_cache_and_chartsheets(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = Workbook()
            sheet = book.active
            sheet["A1"] = 100
            sheet["C1"] = 2
            sheet["B1"] = "=_xll.UNKNOWN(C1)"
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=1, min_row=1))
            book.create_chartsheet("Chart").add_chart(chart)
            book.save(root / "source.xlsx")
            # The add-in result is unavailable locally; its unrelated cached
            # value must survive the input patch.
            with zipfile.ZipFile(root / "source.xlsx") as source, zipfile.ZipFile(root / "cached.xlsx", "w") as target:
                for entry in source.infolist():
                    data = source.read(entry.filename)
                    if entry.filename == "xl/worksheets/sheet1.xml":
                        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                        xml = ElementTree.fromstring(data)
                        cell = xml.find(f".//{{{ns}}}c[@r='B1']")
                        value = cell.find(f"{{{ns}}}v")
                        if value is None:
                            value = ElementTree.SubElement(cell, f"{{{ns}}}v")
                        value.text = "4"
                        data = ElementTree.tostring(xml)
                    target.writestr(entry, data)
            (root / "plan.json").write_text(json.dumps({"changes": [{"sheet": "Sheet", "cell": "A1", "expected_value": 100, "new_value": 120}]}))
            apply_changes(root / "cached.xlsx", root / "plan.json", root / "result.xlsx", root / "result.json")
            result = load_workbook(root / "result.xlsx", data_only=True)
            self.assertEqual(result["Sheet"]["A1"].value, 120)
            self.assertEqual(result["Sheet"]["B1"].value, 4)
            self.assertEqual(result.sheetnames, ["Sheet", "Chart"])
            result.close()


if __name__ == "__main__":
    unittest.main()
