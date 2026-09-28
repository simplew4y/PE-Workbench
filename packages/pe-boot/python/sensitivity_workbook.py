#!/usr/bin/env python3
"""Prepare isolated sensitivity scenarios and read recalculated workbook values."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError("Sensitivity plan must be a JSON object")
    return value


def workbook_options(path: Path, *, data_only: bool) -> dict[str, Any]:
    return {
        "data_only": data_only,
        "keep_links": True,
        "keep_vba": path.suffix.lower() == ".xlsm",
        "read_only": data_only,
    }


def prepare(source: Path, plan_path: Path, output_dir: Path) -> None:
    plan = load_json(plan_path)
    scenarios = plan.get("scenarios")
    if not isinstance(scenarios, list) or not scenarios:
        raise ValueError("Sensitivity plan must contain scenarios")
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source, **workbook_options(source, data_only=False))
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        original_values: dict[tuple[str, str], Any] = {}
        for scenario in scenarios:
            if not isinstance(scenario, dict):
                raise ValueError("Scenario must be an object")
            overrides = scenario.get("overrides", [])
            if not isinstance(overrides, list):
                raise ValueError("Scenario overrides must be a list")
            touched: list[tuple[str, str]] = []
            for override in overrides:
                if not isinstance(override, dict):
                    raise ValueError("Scenario override must be an object")
                sheet_name = str(override["sheet"])
                cell_ref = str(override["cell"])
                key = (sheet_name, cell_ref)
                worksheet = workbook[sheet_name]
                if key not in original_values:
                    original_values[key] = worksheet[cell_ref].value
                if isinstance(original_values[key], str) and original_values[key].startswith("="):
                    raise ValueError(f"Refusing to overwrite formula cell {sheet_name}!{cell_ref}")
                worksheet[cell_ref] = override["value"]
                touched.append(key)

            filename = str(scenario["input_filename"])
            destination = output_dir / filename
            if destination.parent != output_dir or destination.suffix.lower() not in {".xlsx", ".xlsm"}:
                raise ValueError("Invalid scenario filename")
            workbook.save(destination)

            for sheet_name, cell_ref in touched:
                workbook[sheet_name][cell_ref] = original_values[(sheet_name, cell_ref)]
    finally:
        workbook.close()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_results(plan_path: Path, recalculated_dir: Path, output_path: Path) -> None:
    plan = load_json(plan_path)
    observations = plan.get("observations")
    scenarios = plan.get("scenarios")
    if not isinstance(observations, list) or not isinstance(scenarios, list):
        raise ValueError("Sensitivity plan is missing observations or scenarios")

    results: list[dict[str, Any]] = []
    for scenario in scenarios:
        if not isinstance(scenario, dict):
            raise ValueError("Scenario must be an object")
        result_filename = str(scenario["result_filename"])
        result_path = recalculated_dir / result_filename
        if result_path.parent != recalculated_dir or not result_path.is_file():
            raise FileNotFoundError(f"Recalculated scenario is missing: {result_filename}")
        workbook = load_workbook(result_path, **workbook_options(result_path, data_only=True))
        try:
            values: dict[str, Any] = {}
            for observation in observations:
                if not isinstance(observation, dict):
                    raise ValueError("Observation must be an object")
                sheet_name = str(observation["sheet"])
                cell_ref = str(observation["cell"])
                key = str(observation["key"])
                values[key] = json_value(workbook[sheet_name][cell_ref].value)
            results.append({"scenario_id": str(scenario["scenario_id"]), "values": values})
        finally:
            workbook.close()

    with output_path.open("w", encoding="utf-8") as handle:
        json.dump({"scenarios": results}, handle, ensure_ascii=False, separators=(",", ":"))


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare_parser = subparsers.add_parser("prepare")
    prepare_parser.add_argument("--input", required=True, type=Path)
    prepare_parser.add_argument("--plan", required=True, type=Path)
    prepare_parser.add_argument("--output-dir", required=True, type=Path)

    read_parser = subparsers.add_parser("read")
    read_parser.add_argument("--plan", required=True, type=Path)
    read_parser.add_argument("--recalculated-dir", required=True, type=Path)
    read_parser.add_argument("--output", required=True, type=Path)

    args = parser.parse_args()
    if args.command == "prepare":
        prepare(args.input, args.plan, args.output_dir)
    else:
        read_results(args.plan, args.recalculated_dir, args.output)


if __name__ == "__main__":
    main()
