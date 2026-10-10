"""Verify pinned dependencies and real PDF/Excel extraction before deployment."""

import argparse
from importlib.metadata import version
from pathlib import Path
import subprocess
import sys
import tempfile

import pymupdf
from openpyxl import Workbook, load_workbook


def verify(extractor: Path | None = None) -> None:
    requirements = Path(__file__).with_name("requirements.txt")
    for requirement in requirements.read_text(encoding="utf-8").splitlines():
        if not requirement.strip() or requirement.startswith("#"):
            continue
        name, expected = requirement.strip().split("==")
        actual = version(name)
        if actual != expected:
            raise RuntimeError(f"{name}: expected {expected}, installed {actual}")

    with tempfile.TemporaryDirectory(prefix="pe-python-check-") as temporary:
        root = Path(temporary)
        pdf = root / "attachment.pdf"
        xlsx = root / "attachment.xlsx"
        marker = "PE attachment environment ready"
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), marker)
            document.save(pdf)
        workbook = Workbook()
        workbook.active.title = "Model"
        workbook.active.append([marker, 75, "=B1*2"])
        workbook.save(xlsx)
        workbook.close()
        with pymupdf.open(pdf) as document:
            if marker not in document[0].get_text():
                raise RuntimeError("PDF text extraction failed")
        workbook = load_workbook(xlsx, read_only=True, data_only=False)
        try:
            if workbook["Model"]["C1"].value != "=B1*2":
                raise RuntimeError("Excel formula extraction failed")
        finally:
            workbook.close()
        if extractor is not None:
            for source in (pdf, xlsx):
                output = source.with_suffix(".txt")
                subprocess.run([sys.executable, str(extractor), str(source), str(output)],
                               check=True, timeout=30, capture_output=True, text=True)
                extracted = output.read_text(encoding="utf-8")
                if marker not in extracted or (source == xlsx and "=B1*2" not in extracted):
                    raise RuntimeError(f"Attachment extractor failed for {source.suffix}")
    print(f"Shared Python ready: {sys.executable}; pinned dependencies and PDF/Excel checks passed"
          + ("; Web attachment extractor passed" if extractor else ""))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--extractor", type=Path)
    verify(parser.parse_args().extractor)
