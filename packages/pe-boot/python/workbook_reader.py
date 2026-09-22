#!/usr/bin/env python3
"""Read source workbook facts on demand; never calculate or infer financial meaning."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict
from datetime import date, datetime
import hashlib
from heapq import merge
import json
from pathlib import Path
import posixpath
import re
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName

from excel_formula_parser import extract_formula_references

READER_VERSION = "1"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def sha256_text(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def json_safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    return {"type": type(value).__name__, "attributes": json_safe(vars(value))}


def dumps_json(value):
    return json.dumps(json_safe(value), ensure_ascii=False, sort_keys=True, allow_nan=False)


def _formula_details(value, data_type):
    if data_type != "f":
        return False, None, None, {}
    if isinstance(value, str):
        return True, "standard", value, {}
    metadata = json_safe(vars(value))
    formula_type = "array" if type(value).__name__ == "ArrayFormula" else "data_table"
    return True, formula_type, getattr(value, "text", None) or dumps_json(metadata), metadata


def _formula_cache_status(is_formula, value, data_type):
    if not is_formula:
        return "not_applicable"
    if value is None and data_type not in {"str", "s", "inlineStr"}:
        return "missing"
    return "error" if data_type == "e" else "present"


def _relationships(archive, part):
    location = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if location not in archive.namelist():
        return {}
    result = {}
    for item in ET.fromstring(archive.read(location)):
        target = item.get("Target", "")
        result[item.get("Id")] = {
            "relationship_id": item.get("Id"), "type": item.get("Type"),
            "target": target, "external": item.get("TargetMode") == "External",
            "part": target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), target)),
        }
    return result


def _range(min_col, min_row, max_col, max_row):
    first, last = f"{get_column_letter(min_col)}{min_row}", f"{get_column_letter(max_col)}{max_row}"
    return first if first == last else f"{first}:{last}"


def _bounds(value):
    if not isinstance(value, str) or not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9]\d*(?::\$?[A-Za-z]{1,3}\$?[1-9]\d*)?", value):
        raise ValueError("A finite A1 cell or range is required")
    bounds = range_boundaries(value.upper())
    c1, r1, c2, r2 = bounds
    if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576):
        raise ValueError("Range is outside worksheet bounds")
    return bounds


def _integer(request, key, default, minimum, maximum):
    value = request.get(key, default)
    if isinstance(value, bool) or not isinstance(value, int) or not minimum <= value <= maximum:
        raise ValueError(f"{key} must be an integer between {minimum} and {maximum}")
    return value


def _metadata(path, count_cells=False):
    """Stream worksheet XML for structure, including facts unavailable in read-only mode."""
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = _relationships(archive, "xl/workbook.xml")
        sheets = []
        for index, item in enumerate(workbook.findall(f"{NS}sheets/{NS}sheet"), start=1):
            relation = relationships[item.get(REL + "id")]
            if not relation["type"].endswith("/worksheet"):
                continue
            sheet = {
                "sheet_name": item.get("name"), "sheet_index": index,
                "sheet_state": item.get("state", "visible"), "used_range": "",
                "non_empty_cell_count": 0, "formula_count": 0, "error_cell_count": 0, "row_count": 0, "col_count": 0,
                "formula_cache_status_counts": {}, "content_ranges": [],
                "metadata": {"hidden_rows": [], "hidden_columns": [], "merged_ranges": [], "tables": [], "drawings": []},
            }
            metadata = sheet["metadata"]
            cache_counts = Counter()
            min_col = min_row = None
            max_col = max_row = 0
            # Store exact occupied row spans, coalescing identical adjacent spans.
            # These rectangles describe positions only, never business table boundaries.
            active = {}
            regions = []
            sheet_relations = _relationships(archive, relation["part"])
            metadata["comments_part"] = next((r["part"] for r in sheet_relations.values() if r["type"].endswith("/comments")), None)
            metadata["comment_cells"] = sorted(_sheet_comments(archive, sheet), key=lambda ref: (_bounds(ref)[1], _bounds(ref)[0]))
            metadata["conditional_format_ranges"] = []
            with archive.open(relation["part"]) as stream:
                for event, element in ET.iterparse(stream, events=("start", "end")):
                    tag = element.tag.split("}")[-1]
                    if event == "start":
                        if tag == "sheetData":
                            sheet_data = element
                        if tag == "row":
                            row_number = int(element.get("r", "0"))
                            row_columns = []
                            if element.get("hidden") in {"1", "true"}:
                                metadata["hidden_rows"].append(row_number)
                        continue
                    if tag == "c":
                        if count_cells:
                            formula, value, inline = element.find(NS + "f"), element.find(NS + "v"), element.find(NS + "is")
                            if formula is not None or (value is not None and value.text is not None) or inline is not None:
                                c, r, _, _ = _bounds(element.get("r"))
                                row_columns.append(c)
                                min_col, min_row = min(c, min_col or c), min(r, min_row or r)
                                max_col, max_row = max(c, max_col), max(r, max_row)
                                sheet["non_empty_cell_count"] += 1
                                sheet["error_cell_count"] += int(element.get("t") == "e")
                                if formula is not None:
                                    sheet["formula_count"] += 1
                                    kind = element.get("t", "n")
                                    status = "error" if kind == "e" else "present" if (value is not None and value.text is not None) or kind in {"str", "s", "inlineStr"} else "missing"
                                    cache_counts[status] += 1
                        element.clear()
                    elif tag == "row":
                        if count_cells:
                            spans = []
                            for col in sorted(set(row_columns)):
                                if spans and col == spans[-1][1] + 1:
                                    spans[-1] = (spans[-1][0], col)
                                else:
                                    spans.append((col, col))
                            for span, rectangle in list(active.items()):
                                if span not in spans or rectangle[3] != row_number - 1:
                                    regions.append(rectangle)
                                    del active[span]
                            for span in spans:
                                if span in active:
                                    active[span][3] = row_number
                                else:
                                    active[span] = [span[0], row_number, span[1], row_number]
                        element.clear()
                        sheet_data.clear()
                    elif tag == "col":
                        if element.get("hidden") in {"1", "true"}:
                            metadata["hidden_columns"].append({"min": int(element.get("min")), "max": int(element.get("max"))})
                        element.clear()
                    elif tag == "mergeCell":
                        metadata["merged_ranges"].append(element.get("ref"))
                        element.clear()
                    elif tag == "conditionalFormatting":
                        metadata["conditional_format_ranges"].append(element.get("sqref"))
                        element.clear()
                    elif tag == "pane":
                        metadata["pane"] = dict(element.attrib)
                        element.clear()
                    elif tag in {"drawing", "legacyDrawing", "tablePart"}:
                        target = sheet_relations.get(element.get(REL + "id"))
                        if target:
                            if tag == "tablePart":
                                table = ET.fromstring(archive.read(target["part"]))
                                metadata["tables"].append({"name": table.get("name"), "display_name": table.get("displayName"), "range": table.get("ref")})
                            else:
                                metadata["drawings"].append(target)
                        element.clear()
            if count_cells:
                regions.extend(active.values())
                sheet["content_ranges"] = [_range(*rect) for rect in sorted(regions, key=lambda rect: (rect[1], rect[0]))]
                sheet["formula_cache_status_counts"] = dict(cache_counts)
                if min_col is not None:
                    sheet.update(used_range=_range(min_col, min_row, max_col, max_row), row_count=max_row - min_row + 1, col_count=max_col - min_col + 1)
            sheets.append(sheet)
        names = []
        all_sheet_names = [sheet.get("name") for sheet in workbook.findall(f"{NS}sheets/{NS}sheet")]
        for element in workbook.findall(f"{NS}definedNames/{NS}definedName"):
            defined = DefinedName.from_tree(element)
            try:
                destinations = list(defined.destinations)
            except (AttributeError, TypeError, ValueError):
                destinations = []
            local_id = defined.localSheetId
            names.append({
                "name": defined.name, "scope_sheet": all_sheet_names[local_id] if local_id is not None and 0 <= local_id < len(all_sheet_names) else None,
                "name_type": defined.type, "attr_text": defined.attr_text, "hidden": bool(defined.hidden),
                "destinations": destinations, "attributes": dict(element.attrib),
            })
        external_links = []
        for index, external in enumerate(workbook.findall(f"{NS}externalReferences/{NS}externalReference"), start=1):
            target = relationships.get(external.get(REL + "id"))
            if target:
                external_links.append({"index": index, "relationship": target, "sources": list(_relationships(archive, target["part"]).values())})
        calculation = workbook.find(NS + "calcPr")
        return {"reader_version": READER_VERSION, "sheets": sheets, "defined_names": names, "external_links": external_links, "calculation": dict(calculation.attrib) if calculation is not None else {}}


def inspect_workbook(path):
    result = _metadata(path, count_cells=True)
    cache_counts = Counter()
    for sheet in result["sheets"]:
        cache_counts.update(sheet["formula_cache_status_counts"])
    result.update(sheet_count=len(result["sheets"]), formula_count=sum(s["formula_count"] for s in result["sheets"]), non_empty_cell_count=sum(s["non_empty_cell_count"] for s in result["sheets"]), error_cell_count=sum(s["error_cell_count"] for s in result["sheets"]), formula_cache_status_counts=dict(cache_counts), external_link_count=len(result["external_links"]), scan_complete=True)
    return result


def _inspect_page(navigation, request):
    offset, limit = _integer(request, "offset", 0, 0, 2**53 - 1), _integer(request, "limit", 50, 1, 200)
    sheet_name = request.get("sheet")
    sheet = next((item for item in navigation["sheets"] if item["sheet_name"] == sheet_name), None)
    if sheet_name is not None and sheet is None:
        raise ValueError("Unknown sheet")
    section = request.get("section", "content_ranges" if sheet else "sheets")
    summary_keys = ("sheet_name", "sheet_index", "sheet_state", "used_range", "row_count", "col_count", "non_empty_cell_count", "formula_count", "error_cell_count", "formula_cache_status_counts")
    summaries = []
    for item in navigation["sheets"]:
        summaries.append({**{key: item[key] for key in summary_keys}, "content_range_count": len(item["content_ranges"]), **{f"{key}_count": len(value) for key, value in item["metadata"].items() if isinstance(value, list)}, "details_available": True})
    if section == "sheets":
        items = summaries
    elif section in {"defined_names", "external_links"}:
        items = navigation[section]
        if section == "defined_names" and sheet:
            items = [item for item in items if item["scope_sheet"] in {None, sheet_name}]
    elif section in {"content_ranges", "hidden_rows", "hidden_columns", "merged_ranges", "tables", "drawings", "comment_cells"}:
        if sheet is None:
            raise ValueError(f"sheet is required for section {section}")
        items = sheet["content_ranges"] if section == "content_ranges" else sheet["metadata"][section]
    else:
        raise ValueError("Unsupported inspection section")
    page = items[offset:offset + limit]
    next_offset = offset + len(page) if offset + len(page) < len(items) else None
    result = {key: value for key, value in navigation.items() if key not in {"sheets", "defined_names", "external_links"}}
    result.update(section=section, offset=offset, matching_item_count=len(items), next_offset=next_offset, complete=next_offset is None, defined_name_count=len(navigation["defined_names"]))
    result[section] = page
    if sheet:
        result["sheet"] = next(item for item in summaries if item["sheet_name"] == sheet_name)
        result["sheet"]["pane"] = sheet["metadata"].get("pane")
    return result


def _color(color):
    if color is None:
        return None
    return {"type": color.type, "value": color.value, "tint": color.tint}


def _sheet_comments(archive, sheet):
    part = sheet["metadata"].get("comments_part")
    if not part:
        return {}
    root = ET.fromstring(archive.read(part))
    authors = [element.text or "" for element in root.findall(f"{NS}authors/{NS}author")]
    return {
        item.get("ref"): {
            "author": authors[int(item.get("authorId", "0"))],
            "text": "".join(element.text or "" for element in item.findall(f"{NS}text//{NS}t")),
        }
        for item in root.findall(f"{NS}commentList/{NS}comment")
    }


def _cell_record(cell, cached, sheet, comments):
    is_formula, formula_type, formula, formula_metadata = _formula_details(cell.value, cell.data_type)
    status = _formula_cache_status(is_formula, cached.value, cached.data_type)
    cached_value = "" if is_formula and cached.value is None and status == "present" else cached.value
    raw_value = "" if cell.value is None and cell.data_type in {"str", "s", "inlineStr"} else cell.value
    value = json_safe(raw_value)
    saved = json_safe(cached_value) if is_formula else None
    display = saved if is_formula and status in {"present", "error"} else value
    numeric = cached_value if is_formula else cell.value
    metadata = sheet["metadata"]
    merged_range = next((area for area in metadata["merged_ranges"] if _contains(_bounds(area), cell.column, cell.row)), None)
    return {
        "sheet_name": sheet["sheet_name"], "cell_ref": cell.coordinate, "row_index": cell.row, "col_index": cell.column,
        "value_type": f"formula_{formula_type}" if is_formula else type(raw_value).__name__,
        "value": value, "cached": saved,
        "display_value": display if isinstance(display, str) else dumps_json(display) if display is not None else "",
        "raw_value": raw_value if raw_value is None or isinstance(raw_value, str) else dumps_json(raw_value),
        "numeric_value": numeric if isinstance(numeric, (int, float)) and not isinstance(numeric, bool) else None,
        "formula": formula, "cached_value": (cached_value if isinstance(cached_value, str) else dumps_json(cached_value)) if is_formula and cached_value is not None else None,
        "number_format": cell.number_format, "row_label": "", "col_label": "", "period": "", "unit": "",
        "is_formula": int(is_formula), "formula_type": formula_type, "formula_cache_status": status,
        "style": {
            "font_color": _color(cell.font.color), "bold": cell.font.b, "italic": cell.font.i,
            "fill_type": getattr(cell.fill, "patternType", None) or getattr(cell.fill, "type", None),
            "fill_foreground": _color(getattr(cell.fill, "fgColor", None)),
            "fill_background": _color(getattr(cell.fill, "bgColor", None)),
            "fill_gradient": [{"position": stop.position, "color": _color(stop.color)} for stop in getattr(cell.fill, "stop", [])],
        },
        "comment": comments.get(cell.coordinate),
        "metadata_json": dumps_json({
            "sheet_state": sheet["sheet_state"], "hidden_row": cell.row in metadata["hidden_rows"],
            "hidden_column": any(area["min"] <= cell.column <= area["max"] for area in metadata["hidden_columns"]),
            "merged_range": merged_range, "formula_metadata": formula_metadata,
            "conditional_formatting": any(_contains(_bounds(area), cell.column, cell.row)
                for ranges in metadata.get("conditional_format_ranges", []) for area in ranges.split()),
        }),
    }


def _contains(bounds, col, row):
    return bounds[0] <= col <= bounds[2] and bounds[1] <= row <= bounds[3]


def _requests(request, sheet_names):
    ranges = request.get("ranges")
    if ranges is None:
        ranges = [{"sheet": request.get("sheet"), "range": request.get("range")}]
    if not isinstance(ranges, list) or not 1 <= len(ranges) <= 2000:
        raise ValueError("ranges must contain between 1 and 2000 entries")
    result = []
    for area in ranges:
        if not isinstance(area, dict) or area.get("sheet") not in sheet_names:
            raise ValueError("Unknown or missing sheet")
        bounds = _bounds(area.get("range"))
        result.append({"sheet": area["sheet"], "range": _range(*bounds), "bounds": bounds})
    return result


def _source_cells(formula_sheet, value_sheet, comments, kwargs):
    # A comment can point to an empty cell absent from sheetData entirely.
    source = ((cell, cached)
        for row, saved in zip(formula_sheet.iter_rows(**kwargs), value_sheet.iter_rows(**kwargs))
        for cell, cached in zip(row, saved) if hasattr(cell, "coordinate"))
    positions = sorted((_bounds(ref)[1], _bounds(ref)[0]) for ref in comments)
    annotated = ((ReadOnlyCell(formula_sheet, row, col, None), ReadOnlyCell(value_sheet, row, col, None))
        for row, col in positions if kwargs.get("min_row", 1) <= row <= kwargs.get("max_row", 1048576))
    previous = None
    for cell, cached in merge(source, annotated, key=lambda pair: (pair[0].row, pair[0].column)):
        if cell.coordinate != previous:
            yield cell, cached
        previous = cell.coordinate


def _matching_fields(text, comment, query):
    return [field for field, value in (("value", text), ("comment", comment.get("text") if comment else None))
        if isinstance(value, str) and query.casefold() in value.casefold()]


def _read_loaded(formulas, values, metadata, request, action):
    offset, limit = _integer(request, "offset", 0, 0, 2**53 - 1), _integer(request, "limit", 200, 1, 2000)
    names = [sheet["sheet_name"] for sheet in metadata["sheets"]]
    if action == "read":
        ranges = _requests(request, names)
    else:
        query = request.get("query")
        if not isinstance(query, str) or not query:
            raise ValueError("A nonempty search query is required")
        if request.get("sheet") is not None and request["sheet"] not in names:
            raise ValueError("Unknown sheet")
        ranges = []
        query = query.casefold()
    cells, count, non_empty_count = [], 0, 0
    for sheet in metadata["sheets"]:
        name = sheet["sheet_name"]
        selections = [area["bounds"] for area in ranges if area["sheet"] == name]
        selected_cells = {(b[0], b[1]) for b in selections} if selections and all(b[0] == b[2] and b[1] == b[3] for b in selections) else None
        if (action == "read" and not selections) or (action == "search" and request.get("sheet") not in {None, name}):
            continue
        formula_sheet, value_sheet = formulas[name], values[name]
        comments = _sheet_comments(formulas._archive, sheet)
        # ponytail: query scans source XML; add a sheet index only if measured latency warrants it.
        # Never trust worksheet dimensions: third-party producers can save incorrect bounds.
        formula_sheet.reset_dimensions()
        value_sheet.reset_dimensions()
        kwargs = {}
        if selections:
            # Keep each row's actual width; a request for A:XFD must not manufacture 16K empty cells per row.
            kwargs = {"min_row": min(b[1] for b in selections), "max_row": max(b[3] for b in selections)}
        for cell, cached in _source_cells(formula_sheet, value_sheet, comments, kwargs):
            if cell.value is None and cell.data_type not in {"str", "s", "inlineStr"} and cell.coordinate not in comments and not (action == "read" and cell.has_style):
                continue
            if action == "read" and ((cell.column, cell.row) not in selected_cells if selected_cells is not None else not any(_contains(area, cell.column, cell.row) for area in selections)):
                continue
            if action == "search":
                fields = _matching_fields(cached.value if cell.data_type == "f" else cell.value, comments.get(cell.coordinate), query)
                if not fields:
                    continue
            if offset <= count < offset + limit:
                record = _cell_record(cell, cached, sheet, comments)
                if action == "search":
                    record["matched_fields"] = fields
                cells.append(record)
            count += 1
            non_empty_count += int(cell.value is not None or cell.data_type in {"str", "s", "inlineStr"})
    next_offset = offset + len(cells) if offset + len(cells) < count else None
    return {"cells": cells, "matching_cell_count": count, "non_empty_cell_count": non_empty_count, "offset": offset, "next_offset": next_offset, "complete": next_offset is None, "scan_complete": True, "requested_ranges": [{"sheet": area["sheet"], "range": area["range"]} for area in ranges]}


def build_text_index(path, source_sha256, navigation=None):
    """Index original text, saved text results and comments at their exact locations."""
    metadata = navigation if navigation is not None else _metadata(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        entries = []
        sheets = []
        for sheet_index, sheet in enumerate(workbook.worksheets):
            sheets.append(sheet.title)
            comments = _sheet_comments(workbook._archive, next(item for item in metadata["sheets"] if item["sheet_name"] == sheet.title))
            text_cells = {ref: {"comment": comment["text"]} for ref, comment in comments.items()}
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        text_cells.setdefault(cell.coordinate, {})["value"] = cell.value
            for ref in sorted(text_cells, key=lambda ref: (_bounds(ref)[1], _bounds(ref)[0])):
                entries.append([sheet_index, ref, text_cells[ref]])
        return {"reader_version": READER_VERSION, "source_sha256": source_sha256, "sheets": sheets, "entries": entries}
    finally:
        workbook.close()


def _search_index(formulas, values, metadata, request, text_index):
    offset, limit = _integer(request, "offset", 0, 0, 2**53 - 1), _integer(request, "limit", 200, 1, 2000)
    query = request.get("query")
    if not isinstance(query, str) or not query:
        raise ValueError("A nonempty search query is required")
    if request.get("sheet") is not None and request["sheet"] not in formulas.sheetnames:
        raise ValueError("Unknown sheet")
    if text_index.get("reader_version") != READER_VERSION or text_index.get("sheets") != [sheet.title for sheet in formulas.worksheets]:
        raise ValueError("Text index does not match workbook reader or sheets")
    count, selected = 0, []
    matches = {}
    for sheet_index, cell_ref, texts in text_index["entries"]:
        sheet_name = text_index["sheets"][sheet_index]
        fields = _matching_fields(texts.get("value"), {"text": texts.get("comment")}, query)
        if request.get("sheet") not in {None, sheet_name} or not fields:
            continue
        if offset <= count < offset + limit:
            selected.append({"sheet": sheet_name, "range": cell_ref})
            matches[(sheet_name, cell_ref)] = fields
        count += 1
    result = _read_loaded(formulas, values, metadata, {"ranges": selected, "limit": limit}, "read") if selected else {"cells": []}
    if len(result["cells"]) != len(selected):
        raise ValueError("Text index contains stale source locations")
    for cell in result["cells"]:
        cell["matched_fields"] = matches[(cell["sheet_name"], cell["cell_ref"])]
    next_offset = offset + len(selected) if offset + len(selected) < count else None
    return {"cells": result["cells"], "matching_cell_count": count, "offset": offset, "next_offset": next_offset, "complete": next_offset is None, "scan_complete": True, "requested_ranges": [], "index_used": True}


def _trace_loaded(formulas, values, metadata, request):
    max_depth, max_nodes = _integer(request, "max_depth", 2, 0, 20), _integer(request, "max_nodes", 100, 1, 2000)
    roots = _requests(request, formulas.sheetnames)
    pending = [{"sheet": area["sheet"], "range": area["range"], "depth": 0} for area in roots]
    nodes, edges, issues, unvisited, pending_reads = {}, [], [], [], []
    names = metadata["defined_names"]
    sheet_names = {name.casefold(): name for name in formulas.sheetnames}
    while pending:
        depth = pending[0]["depth"]
        batch = [area for area in pending if area["depth"] == depth]
        areas = batch[:100]
        pending = batch[100:] + [area for area in pending if area["depth"] != depth]
        if depth > max_depth or len(nodes) >= max_nodes:
            unvisited.extend(areas)
            continue
        result = _read_loaded(formulas, values, metadata, {"ranges": areas, "limit": max_nodes}, "read")
        if not result["complete"]:
            pending_reads.append({"ranges": [{"sheet": area["sheet"], "range": area["range"]} for area in areas], "offset": result["next_offset"], "depth": depth})
        for cell in result["cells"]:
            key = (cell["sheet_name"], cell["cell_ref"])
            if key in nodes:
                continue
            if len(nodes) >= max_nodes:
                unvisited.append({"sheet": key[0], "range": key[1], "depth": depth})
                continue
            cell["depth"] = depth
            nodes[key] = cell
            if not cell["is_formula"]:
                continue
            if cell["formula_cache_status"] in {"missing", "error"}:
                issues.append({"sheet": key[0], "cell_ref": key[1], "reason": "formula_cache_unavailable", "status": cell["formula_cache_status"]})
            formula = cell["formula"] or ""
            if cell["formula_type"] == "data_table" or re.search(r"\b(?:INDIRECT|OFFSET)\s*\(", formula, re.IGNORECASE):
                issues.append({"sheet": key[0], "cell_ref": key[1], "reason": "dynamic_reference", "formula": formula})
            references = extract_formula_references(formula, source_sheet=key[0], defined_names=[item["name"] for item in names])
            for reference in references:
                edge = {"source_sheet": key[0], "source_cell_ref": key[1], **asdict(reference)}
                destinations = []
                if reference.reference_kind == "defined_name":
                    candidates = [item for item in names if item["name"].casefold() == reference.defined_name.casefold() and item["scope_sheet"] in {None, key[0]}]
                    candidates.sort(key=lambda item: item["scope_sheet"] is None)
                    if candidates:
                        destinations = candidates[0]["destinations"]
                        edge["destinations"] = destinations
                        edge["definition"] = candidates[0]["attr_text"]
                        if destinations:
                            edge["parse_status"] = "resolved"
                elif reference.parse_status == "resolved":
                    destinations = [(reference.target_sheet, reference.target_range)]
                destinations = [(sheet_names.get(name.casefold(), name), area) for name, area in destinations]
                if destinations:
                    edge["destinations"] = destinations
                    if reference.target_sheet:
                        edge["target_sheet"] = sheet_names.get(reference.target_sheet.casefold(), reference.target_sheet)
                if reference.external_workbook:
                    edge["external_sources"] = [link for link in metadata["external_links"] if str(link["index"]) == reference.external_workbook]
                edges.append(edge)
                if not destinations:
                    issues.append({"sheet": key[0], "cell_ref": key[1], "reason": edge["parse_status"], "reference": reference.raw_reference})
                for target_sheet, target_range in destinations:
                    try:
                        bounds = _bounds(target_range)
                    except ValueError:
                        issues.append({"sheet": key[0], "cell_ref": key[1], "reason": "non_finite_reference", "reference": reference.raw_reference})
                        continue
                    if target_sheet not in formulas.sheetnames:
                        issues.append({"sheet": key[0], "cell_ref": key[1], "reason": "missing_sheet", "reference": reference.raw_reference})
                        continue
                    if (target_sheet, _range(*bounds)) not in nodes:
                        pending.append({"sheet": target_sheet, "range": _range(*bounds), "depth": depth + 1})
    # Only a back edge in the explored graph is a cycle; shared inputs are normal.
    graph = {key: set() for key in nodes}
    for edge in edges:
        if edge["parse_status"] != "resolved":
            continue
        destinations = edge.get("destinations") or [(edge["target_sheet"], edge["target_range"])]
        for target_sheet, target_range in destinations:
            try:
                bounds = _bounds(target_range)
            except ValueError:
                continue
            graph[(edge["source_sheet"], edge["source_cell_ref"])].update(key for key, cell in nodes.items() if key[0] == target_sheet and _contains(bounds, cell["col_index"], cell["row_index"]))
    completed = set()
    for root in graph:
        if root in completed:
            continue
        path = [root]
        active = {root}
        stack = [iter(sorted(graph[root]))]
        while stack:
            target = next(stack[-1], None)
            if target is None:
                completed.add(path[-1])
                active.remove(path.pop())
                stack.pop()
            elif target in active:
                issues.append({"reason": "circular_reference", "cells": [{"sheet": key[0], "cell_ref": key[1]} for key in path[path.index(target):] + [target]]})
            elif target not in completed:
                path.append(target)
                active.add(target)
                stack.append(iter(sorted(graph[target])))
    return {"nodes": list(nodes.values()), "edges": edges, "issues": issues, "pending_ranges": unvisited, "pending_reads": pending_reads, "complete": all(issue["reason"] == "formula_cache_unavailable" for issue in issues) and not unvisited and not pending_reads, "truncated": bool(unvisited or pending_reads)}


def read_workbook(path, request, text_index_path=None, navigation=None):
    if not isinstance(request, dict):
        raise ValueError("Reader request must be an object")
    action = request.get("action")
    if action == "inspect":
        return _inspect_page(navigation if navigation is not None else inspect_workbook(path), request)
    if action not in {"read", "search", "trace"}:
        raise ValueError("Unsupported reader action")
    metadata = navigation if navigation is not None else _metadata(path)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if action == "trace":
                return _trace_loaded(formulas, values, metadata, request)
            if action == "search" and text_index_path is not None:
                with Path(text_index_path).open(encoding="utf-8") as stream:
                    return _search_index(formulas, values, metadata, request, json.load(stream))
            return _read_loaded(formulas, values, metadata, request, action)
        finally:
            values.close()
    finally:
        formulas.close()


def navigation_artifact(path, dataset_id, doc_id):
    navigation = inspect_workbook(path)
    identity = {"dataset_id": dataset_id, "doc_id": doc_id}
    sheets, regions, names = [], [], []
    for sheet in navigation["sheets"]:
        name = sheet["sheet_name"]
        density = sheet["formula_count"] / max(1, sheet["non_empty_cell_count"])
        sheets.append({
            **identity, **{key: sheet[key] for key in ("sheet_index", "sheet_name", "sheet_state", "used_range", "row_count", "col_count", "non_empty_cell_count", "formula_count")},
            "sheet_id": sha256_text(f"{doc_id}\0sheet\0{name}")[:40], "sheet_role": "worksheet", "formula_density": density,
            "summary": f"{name}: {sheet['used_range'] or 'empty'}; {sheet['non_empty_cell_count']} occupied cells; {sheet['formula_count']} formulas",
            "header_json": "[]", "metadata_json": dumps_json({**sheet["metadata"], "content_ranges": sheet["content_ranges"], "formula_cache_status_counts": sheet["formula_cache_status_counts"]}),
        })
        # Exact content positions live in sheet metadata, avoiding a second copy per rectangle.
    for index, name in enumerate(navigation["defined_names"]):
        names.append({
            **identity, **{key: name[key] for key in ("name", "scope_sheet", "name_type", "attr_text")},
            "defined_name_id": sha256_text(f"{doc_id}\0defined_name\0{name['scope_sheet'] or ''}\0{name['name']}\0{index}")[:40],
            "hidden": int(name["hidden"]), "metadata_json": dumps_json({"attributes": name["attributes"], "destinations": name["destinations"]}),
        })
    cache_counts = Counter()
    for sheet in navigation["sheets"]:
        cache_counts.update(sheet["formula_cache_status_counts"])
    workbook = {
        **identity, "workbook_id": sha256_text(f"{doc_id}\0workbook")[:40], "workbook_type": "workbook",
        "sheet_count": navigation["sheet_count"], "visible_sheet_count": sum(s["sheet_state"] == "visible" for s in sheets),
        "formula_count": navigation["formula_count"], "non_empty_cell_count": navigation["non_empty_cell_count"],
        "formula_density": navigation["formula_count"] / max(1, navigation["non_empty_cell_count"]),
        "metadata_json": dumps_json({"reader_version": READER_VERSION, "calculation": navigation["calculation"], "external_links": navigation["external_links"], "external_link_count": len(navigation["external_links"]), "defined_name_count": len(names), "formula_cache_status_counts": dict(cache_counts), "error_cell_count": navigation["error_cell_count"], "scan_complete": True}),
    }
    warnings = []
    unavailable_count = cache_counts["missing"] + cache_counts["error"]
    if unavailable_count:
        warnings.append(f"文件中有 {unavailable_count} 个公式的已保存结果缺失或包含错误；reader 不重新计算公式。")
    if navigation["external_links"]:
        warnings.append(f"文件包含 {len(navigation['external_links'])} 个外部数据链接；当前读取文件中已保存的结果。")
    return {
        "parser_name": "workbook_reader", "parser_version": READER_VERSION,
        "sheet_count": navigation["sheet_count"], "cell_count": navigation["non_empty_cell_count"], "formula_count": navigation["formula_count"],
        "warnings": warnings, "navigation": navigation,
        "tables": {"excel_workbooks": [workbook], "excel_sheets": sheets, "excel_regions": regions, "excel_defined_names": names, "excel_cells": [], "excel_formula_references": [], "valuation_date_candidates": [], "metric_facts": []},
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--text-index", type=Path)
    parser.add_argument("--navigation", type=Path)
    args = parser.parse_args()
    path = args.input.resolve(strict=True)
    request = json.load(sys.stdin)
    navigation = request.pop("_navigation", None)
    if args.navigation:
        with args.navigation.open(encoding="utf-8") as stream:
            navigation = json.load(stream)["navigation"]
    prepared = navigation is None
    if prepared:
        navigation = inspect_workbook(path)
    result = read_workbook(path, request, args.text_index, navigation)
    # Private transport field: callers retain navigation, never put it in model context.
    if prepared:
        result["_navigation"] = navigation
    json.dump(result, sys.stdout, ensure_ascii=False, allow_nan=False)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
