"""Source-backed reader regression: navigation, pagination, and formula evidence."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Color, Font, GradientFill, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from workbook_reader import NS, REL, build_text_index, inspect_workbook, read_workbook


def make_model(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Valuation"
    sheet.append(["Target Price", "=Named*2", "=B1+Inputs!A1"])
    sheet["B1"].number_format = '"EUR/share" 0.00'
    sheet["E3"] = "layout"
    sheet.merge_cells("E3:G3")
    sheet["A8"] = "=INDIRECT(\"Inputs!A1\")"
    sheet["A9"] = "='[1]Sheet1'!B2"
    sheet["K5"] = "=K6"
    sheet["K6"] = "=K5"
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions.group("E", "G", hidden=True)
    for row in range(20, 40):
        sheet.cell(row, 1, "x" * 30000)
    sheet["A100"] = "TARGET at the end"
    sheet["B100"] = " #REF! is a note "
    sheet["C100"] = "=literal string"
    sheet["C100"].data_type = "s"
    sheet["D100"] = ""
    inputs = workbook.create_sheet("Inputs")
    inputs["A1"] = 12
    inputs["B1"] = 3
    inputs.sheet_state = "veryHidden"
    workbook.defined_names.add(DefinedName("Named", attr_text="'Inputs'!$A$1"))
    sheet.defined_names.add(DefinedName("LocalName", attr_text="'Inputs'!$B$1"))
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    xml = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    xml.find(NS + "dimension").set("ref", "A1:A1")
    for cell in xml.findall(f"{NS}sheetData/{NS}row/{NS}c"):
        if cell.get("r") == "B1":
            cell.find(NS + "v").text = "24"
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
    xml = ET.fromstring(entries["xl/workbook.xml"])
    external = ET.SubElement(xml, NS + "externalReferences")
    ET.SubElement(external, NS + "externalReference", {REL + "id": "rIdExternal"})
    entries["xl/workbook.xml"] = ET.tostring(xml)
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    xml = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
    ET.SubElement(xml, rel_ns + "Relationship", {"Id": "rIdExternal", "Type": REL[1:-1] + "/externalLink", "Target": "externalLinks/externalLink1.xml"})
    entries["xl/_rels/workbook.xml.rels"] = ET.tostring(xml)
    xml = ET.Element(rel_ns + "Relationships")
    ET.SubElement(xml, rel_ns + "Relationship", {"Id": "rId1", "Type": REL[1:-1] + "/externalLinkPath", "Target": "file:///source/prior.xlsx", "TargetMode": "External"})
    entries["xl/externalLinks/_rels/externalLink1.xml.rels"] = ET.tostring(xml)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


class WorkbookReaderTests(unittest.TestCase):
    def test_comment_navigation_and_search_match_with_and_without_index(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Comments.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Model"
            sheet["A1"] = "Margin assumption"
            sheet["A1"].comment = Comment("Margin assumption from analyst", "Analyst")
            sheet["B2"] = 0.25
            sheet["B2"].comment = Comment("Margin assumption is not a reported actual", "Analyst")
            sheet["D80"].comment = Comment("Margin assumption pending review", "Reviewer")
            hidden = book.create_sheet("Hidden")
            hidden.sheet_state = "veryHidden"
            hidden["C10"].comment = Comment("Margin downside", "Reviewer")
            book.save(path)
            book.close()
            # Legal OOXML: annotation can exist without a corresponding cell element.
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            xml = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
            for row in xml.findall(f"{NS}sheetData/{NS}row"):
                for cell in list(row):
                    if cell.get("r") == "D80":
                        row.remove(cell)
            parts["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
            with ZipFile(path, "w", ZIP_DEFLATED) as target:
                for name, data in parts.items():
                    target.writestr(name, data)
            original = path.read_bytes()
            navigation = inspect_workbook(path)
            self.assertEqual(navigation["sheets"][0]["used_range"], "A1:B2")
            overview = read_workbook(path, {"action": "inspect"}, navigation=navigation)
            self.assertEqual([sheet["comment_cells_count"] for sheet in overview["sheets"]], [3, 1])
            request = {"action": "inspect", "section": "comment_cells", "sheet": "Model", "limit": 2}
            page = read_workbook(path, request, navigation=navigation)
            self.assertEqual(page["comment_cells"], ["A1", "B2"])
            self.assertEqual(page["next_offset"], 2)
            self.assertEqual(read_workbook(path, {**request, "offset": 2}, navigation=navigation)["comment_cells"], ["D80"])
            index = Path(directory) / "text-index.json"
            with patch("workbook_reader._metadata", side_effect=AssertionError("Must reuse preparation")):
                index.write_text(json.dumps(build_text_index(path, "test-source", navigation)))
            for offset in range(4):
                query = {"action": "search", "query": "MARGIN", "offset": offset, "limit": 1}
                direct = read_workbook(path, query, navigation=navigation)
                indexed = read_workbook(path, query, index, navigation)
                self.assertEqual(direct["cells"], indexed["cells"])
                self.assertEqual(indexed["matching_cell_count"], 4)
                self.assertEqual(indexed["next_offset"], offset + 1 if offset < 3 else None)
                self.assertEqual(indexed["cells"][0]["matched_fields"], ["value", "comment"] if offset == 0 else ["comment"])
            blank = read_workbook(path, {"action": "read", "sheet": "Model", "range": "D80"}, navigation=navigation)
            self.assertIsNone(blank["cells"][0]["value"])
            self.assertIsNone(blank["cells"][0]["raw_value"])
            self.assertEqual(blank["non_empty_cell_count"], 0)
            self.assertEqual(blank["cells"][0]["comment"]["text"], "Margin assumption pending review")
            validation = {"action": "validate", "ranges": [
                {"sheet": "Model", "range": "D80"},
                {"sheet": "Hidden", "range": "C10"},
                {"sheet": "Model", "range": "D81"},
            ]}
            expected = [{**area, "exists": index < 2} for index, area in enumerate(validation["ranges"])]
            self.assertEqual(read_workbook(path, validation)["ranges"], expected)
            self.assertEqual(read_workbook(path, validation, navigation=navigation)["ranges"], expected)
            filtered = read_workbook(path, {"action": "search", "query": "margin", "sheet": "Hidden"}, index, navigation)
            self.assertEqual(filtered["matching_cell_count"], 1)
            self.assertEqual(path.read_bytes(), original)

    def test_styles_comments_and_prepared_navigation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Styles.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Inputs"
            sheet["A1"] = 47
            sheet["A1"].font = Font(color="FF0000FF", bold=True)
            sheet["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
            sheet["A1"].comment = Comment("EUR per share\nAnalyst input", "Analyst")
            sheet["B1"].font = Font(color=Color(theme=4, tint=0.25))
            sheet["C1"].fill = PatternFill("solid", fgColor=Color(indexed=3))
            sheet["D1"].comment = Comment("Blank input, not zero", "Reviewer")
            sheet["D1"].fill = GradientFill(stop=["FF0000FF", "FFFFFFFF"])
            sheet.conditional_formatting.add("A1:D1", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFFF0000")))
            workbook.save(path)
            workbook.close()
            navigation = inspect_workbook(path)
            with patch("workbook_reader._metadata", side_effect=AssertionError("Must reuse preparation")):
                overview = read_workbook(path, {"action": "inspect"}, navigation=navigation)
                result = read_workbook(path, {"action": "read", "sheet": "Inputs", "range": "A1:D1"}, navigation=navigation)
            self.assertEqual(overview["sheet_count"], 1)
            cells = result["cells"]
            self.assertEqual(len(cells), 4)
            self.assertEqual(result["non_empty_cell_count"], 1)
            self.assertEqual(cells[0]["style"]["font_color"], {"type": "rgb", "value": "FF0000FF", "tint": 0.0})
            self.assertEqual(cells[0]["style"]["fill_foreground"]["value"], "FFFFFF00")
            self.assertTrue(cells[0]["style"]["bold"])
            self.assertEqual(cells[0]["comment"]["text"], "EUR per share\nAnalyst input")
            self.assertEqual(cells[0]["comment"]["author"], "Analyst")
            self.assertTrue(json.loads(cells[0]["metadata_json"])["conditional_formatting"])
            self.assertEqual(cells[1]["style"]["font_color"], {"type": "theme", "value": 4, "tint": 0.25})
            self.assertEqual(cells[2]["style"]["fill_foreground"]["type"], "indexed")
            self.assertIsNone(cells[3]["value"])
            self.assertEqual(cells[3]["comment"]["text"], "Blank input, not zero")
            self.assertEqual(cells[3]["style"]["fill_gradient"][0]["color"]["value"], "FF0000FF")

    def test_batch_validation_matches_read_existence_without_recalculating_or_scanning_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Validation.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Model"
            sheet["A1"] = 0
            sheet["B1"] = ""
            sheet["C1"] = "=Missing!A1"
            sheet["D1"].number_format = "0.00"
            sheet["A4"].comment = Comment("Pending", "Reader")
            other = book.create_sheet("Other")
            other["D20"] = False
            book.save(path)
            book.close()
            ranges = [{"sheet": "Model", "range": area} for area in ["A1", "B1", "C1", "D1", "E1", "A4", "A1:E9", "XFD100"]]
            ranges.append({"sheet": "Other", "range": "D20"})
            expected = [{**area, "exists": bool(read_workbook(path, {"action": "read", **area})["cells"])} for area in ranges]
            with patch("workbook_reader.load_workbook", wraps=load_workbook) as loader, patch("workbook_reader._metadata", side_effect=AssertionError("Validation must not scan whole-workbook metadata")):
                result = read_workbook(path, {"action": "validate", "ranges": ranges})
            self.assertEqual(result["ranges"], expected)
            self.assertEqual(loader.call_count, 1)
            self.assertTrue(loader.call_args.kwargs["read_only"])
            self.assertFalse(loader.call_args.kwargs["data_only"])
            with self.assertRaisesRegex(ValueError, "Unknown or missing sheet"):
                read_workbook(path, {"action": "validate", "ranges": [{"sheet": "Missing", "range": "A1"}]})

    def test_case_insensitive_formula_destinations_and_cycles(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Case.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Output"
            sheet["A1"] = "=inputs!A1+Named"
            sheet["A2"] = "=inputs!A2"
            inputs = workbook.create_sheet("Inputs")
            inputs["A1"] = 47
            inputs["A2"] = "=OUTPUT!A2"
            workbook.defined_names.add(DefinedName("Named", attr_text="'INPUTS'!$A$1"))
            workbook.save(path)
            workbook.close()
            trace = read_workbook(path, {"action": "trace", "sheet": "Output", "range": "A1"})
            self.assertTrue(trace["complete"], trace)
            self.assertEqual([(n["sheet_name"], n["cell_ref"]) for n in trace["nodes"]], [("Output", "A1"), ("Inputs", "A1")])
            self.assertTrue(all(edge["destinations"][0][0] == "Inputs" for edge in trace["edges"]))
            cycle = read_workbook(path, {"action": "trace", "sheet": "Output", "range": "A2", "max_depth": 4})
            self.assertIn("circular_reference", [issue["reason"] for issue in cycle["issues"]])
            self.assertNotIn("missing_sheet", [issue["reason"] for issue in cycle["issues"]])

    def test_source_navigation_read_search_trace_and_lightweight_cli(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Model.xlsx"
            make_model(path)
            before = path.read_bytes()
            navigation = inspect_workbook(path)
            sheet = navigation["sheets"][0]
            self.assertEqual(sheet["used_range"], "A1:K100")
            self.assertNotIn("A1:K100", sheet["content_ranges"])
            self.assertIn("E3", sheet["content_ranges"])
            self.assertEqual(sheet["metadata"]["hidden_columns"], [{"min": 5, "max": 7}])
            self.assertEqual(sheet["metadata"]["hidden_rows"], [3])
            self.assertEqual(navigation["sheets"][1]["sheet_state"], "veryHidden")
            self.assertEqual(navigation["formula_cache_status_counts"], {"present": 1, "missing": 5})
            self.assertEqual(navigation["external_links"][0]["sources"][0]["target"], "file:///source/prior.xlsx")
            self.assertEqual(next(n for n in navigation["defined_names"] if n["name"] == "LocalName")["scope_sheet"], "Valuation")
            overview = read_workbook(path, {"action": "inspect", "limit": 1})
            self.assertNotIn("content_ranges", overview["sheets"][0])
            self.assertEqual(overview["next_offset"], 1)
            details = read_workbook(path, {"action": "inspect", "sheet": "Valuation", "limit": 1})
            self.assertEqual(details["section"], "content_ranges")
            self.assertEqual(details["content_ranges"], ["A1:C1"])
            self.assertFalse(details["complete"])
            named_page = read_workbook(path, {"action": "inspect", "section": "defined_names", "limit": 1})
            self.assertEqual(named_page["matching_item_count"], 2)
            request = {"action": "read", "sheet": "Valuation", "range": "A1:C100", "limit": 2}
            page = read_workbook(path, request)
            self.assertFalse(page["complete"])
            self.assertEqual(page["next_offset"], 2)
            formula = page["cells"][1]
            self.assertEqual(formula["formula"], "=Named*2")
            self.assertEqual(formula["cached"], 24)
            self.assertEqual(formula["numeric_value"], 24)
            self.assertEqual(formula["unit"], "")
            page2 = read_workbook(path, {**request, "offset": page["next_offset"]})
            self.assertEqual(page2["cells"][0]["cell_ref"], "C1")
            late = read_workbook(path, {"action": "search", "query": "target", "limit": 1})
            self.assertEqual(late["matching_cell_count"], 2)
            late = read_workbook(path, {"action": "search", "query": "target", "offset": late["next_offset"], "limit": 1})
            self.assertEqual(late["cells"][0]["cell_ref"], "A100")
            self.assertTrue(late["complete"])
            self.assertTrue(late["scan_complete"])
            long = read_workbook(path, {"action": "read", "sheet": "Valuation", "range": "A20"})
            self.assertEqual(len(long["cells"][0]["raw_value"]), 30000)
            literal = read_workbook(path, {"action": "read", "sheet": "Valuation", "range": "C100"})
            self.assertEqual(literal["cells"][0]["value"], "=literal string")
            self.assertFalse(literal["cells"][0]["is_formula"])
            empty_text = read_workbook(path, {"action": "read", "sheet": "Valuation", "range": "D100"})
            self.assertEqual(empty_text["cells"][0]["value"], "")
            merged = read_workbook(path, {"action": "read", "sheet": "Valuation", "range": "E3"})
            self.assertEqual(json.loads(merged["cells"][0]["metadata_json"])["merged_range"], "E3:G3")
            trace = read_workbook(path, {"action": "trace", "sheet": "Valuation", "range": "C1", "max_depth": 3})
            self.assertTrue(trace["complete"], trace)
            self.assertEqual({(c["sheet_name"], c["cell_ref"]) for c in trace["nodes"]}, {("Valuation", "C1"), ("Valuation", "B1"), ("Inputs", "A1")})
            limited = read_workbook(path, {"action": "trace", "sheet": "Valuation", "range": "C1", "max_depth": 0})
            self.assertTrue(limited["truncated"])
            self.assertFalse(limited["complete"])
            self.assertTrue(limited["pending_ranges"])
            external = read_workbook(path, {"action": "trace", "sheet": "Valuation", "range": "A9"})
            self.assertFalse(external["complete"])
            self.assertEqual(external["edges"][0]["external_sources"][0]["index"], 1)
            dynamic = read_workbook(path, {"action": "trace", "sheet": "Valuation", "range": "A8"})
            self.assertFalse(dynamic["complete"])
            self.assertIn("dynamic_reference", [issue["reason"] for issue in dynamic["issues"]])
            cycle = read_workbook(path, {"action": "trace", "sheet": "Valuation", "range": "K5", "max_depth": 3})
            self.assertFalse(cycle["complete"])
            self.assertIn("circular_reference", [issue["reason"] for issue in cycle["issues"]])
            partial = read_workbook(path, {"action": "trace", "ranges": [{"sheet": "Valuation", "range": "A1"}, {"sheet": "Inputs", "range": "A1:B1"}], "max_nodes": 1})
            self.assertEqual(partial["pending_reads"][0]["offset"], 1)
            remaining = read_workbook(path, {"action": "read", **{key: partial["pending_reads"][0][key] for key in ("ranges", "offset")}})
            self.assertEqual([c["cell_ref"] for c in remaining["cells"]], ["A1", "B1"])
            with self.assertRaises(ValueError):
                read_workbook(path, {"action": "read", "sheet": "Valuation", "range": "A:A"})
            output = Path(directory) / "navigation.json"
            index = Path(directory) / "text-index.json"
            command = [sys.executable, str(Path(__file__).with_name("parse_workbook.py")), "--input", str(path), "--output", str(output), "--text-index-output", str(index), "--doc-id", "a" * 40, "--dataset-id", "dataset", "--revision", "1", "--sha256", hashlib.sha256(before).hexdigest(), "--filename", path.name]
            subprocess.run(command, check=True, capture_output=True, text=True)
            artifact = json.loads(output.read_text())
            self.assertEqual(artifact["parser_name"], "workbook_reader")
            self.assertEqual(len(artifact["warnings"]), 2)
            for table in ("excel_cells", "excel_formula_references", "metric_facts", "valuation_date_candidates"):
                self.assertEqual(artifact["tables"][table], [])
            self.assertLess(output.stat().st_size, 15000)
            indexed = read_workbook(path, {"action": "search", "query": "target", "offset": 1, "limit": 1}, index)
            self.assertTrue(indexed["index_used"])
            self.assertEqual(indexed["cells"], late["cells"])
            self.assertEqual(indexed["matching_cell_count"], 2)
            indexed_literal = read_workbook(path, {"action": "search", "query": "literal"}, index)
            self.assertEqual(indexed_literal["cells"], [{**cell, "matched_fields": ["value"]} for cell in literal["cells"]])
            self.assertEqual(path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
