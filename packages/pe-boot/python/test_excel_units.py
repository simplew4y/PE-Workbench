from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from workbook import parse_workbook


class FormulaUnitsTest(unittest.TestCase):
    def parse(self, workbook: Workbook) -> dict[tuple[str, str], dict]:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'arbitrary.xlsx'
            workbook.save(path)
            result = parse_workbook(dataset_id='test', doc_id='doc_' + 'b' * 24, path=path)
        self.facts = result['tables']['metric_facts']
        return {(cell['sheet_name'], cell['cell_ref']): cell for cell in result['tables']['excel_cells']}

    def test_renamed_shifted_cross_sheet_valuation_and_share_scale(self):
        for shift, col, name, currency in [(0, 2, 'Source', 'EUR'), (14, 16, "Firm's estimates", 'USD')]:
            with self.subTest(name=name):
                workbook = Workbook()
                source = workbook.active
                source.title = name
                source.cell(shift + 1, 1, f'P&L in {currency} million')
                source.cell(shift + 2, 1, 'Net income')
                income = source.cell(shift + 2, col, 600)
                quoted = name.replace("'", "''")
                target = workbook.create_sheet('Results')
                target.append(['Net income', f"='{quoted}'!{income.coordinate}"])
                target.append(['Target multiple', 18])
                target.append(['Implied equity value', '=B1*B2'])
                target.append(['Number of shares', 200])
                target.append(['Fair value per share', '=ROUND(B3/B4,1)'])
                target.append(['Weighted target price', '=AVERAGE(B5,B5)'])
                cells = self.parse(workbook)
                for ref, unit in [('B1', currency + 'm'), ('B3', currency + 'm'), ('B4', 'shares_m'),
                                  ('B5', currency + '/share'), ('B6', currency + '/share')]:
                    self.assertEqual(cells['Results', ref]['unit'], unit, ref)
                self.assertEqual(cells['Results', 'B4']['numeric_value'], 200)
                self.assertEqual(cells['Results', 'B5']['formula_cache_status'], 'missing')
                self.assertIsNone(cells['Results', 'B5']['numeric_value'])
                context = json.loads(cells['Results', 'B4']['metadata_json'])['unit_context']
                self.assertEqual(context['method'], 'formula_lineage')
                self.assertTrue(any(source['sheet_name'] == 'Results' and source['cell_ref'] == 'B5' for source in context['sources']))
                fact = next(fact for fact in self.facts if fact['sheet_name'] == 'Results' and fact['cell_ref'] == 'B4')
                self.assertEqual(fact['unit'], 'shares_m')
                self.assertNotIn('share_count_scale_not_explicit', json.loads(fact['quality_issues_json']))

    def test_resolved_names_ranges_and_reverse_sheet_order(self):
        workbook = Workbook()
        result = workbook.active
        result.title = 'Output first'
        result.append(['Equity value', '=SUM(Amounts)'])
        result.append(['Number of shares (million)', 50])
        result.append(['Target price', '=B1/B2'])
        source = workbook.create_sheet('Inputs later')
        source.append(['人民币百万元'])
        source.append(['Cash', 100])
        source.append(['Cash', 200])
        workbook.defined_names.add(DefinedName('Amounts', attr_text="'Inputs later'!$B$2:$B$3"))
        cells = self.parse(workbook)
        self.assertEqual(cells['Output first', 'B1']['unit'], 'CNYm')
        self.assertEqual(cells['Output first', 'B3']['unit'], 'CNY/share')

    def test_multiple_currencies_are_scoped_to_their_formula_sources(self):
        workbook = Workbook()
        source = workbook.active
        source.title = 'Inputs'
        source.append(['EURm', None, None, 'USDm'])
        source.append(['Equity', 100, None, 'Equity', 200])
        source.append(['Number of shares (million)', 10, None, 'Number of shares (million)', 20])
        output = workbook.create_sheet('Quotes')
        output.append(['Target price', '=Inputs!B2/Inputs!B3'])
        output.append(['Target price', '=Inputs!E2/Inputs!E3'])
        output.append(['Target price', 15])
        cells = self.parse(workbook)
        self.assertEqual(cells['Quotes', 'B1']['unit'], 'EUR/share')
        self.assertEqual(cells['Quotes', 'B2']['unit'], 'USD/share')
        self.assertEqual(cells['Quotes', 'B3']['unit'], 'per_share')

    def test_conflicting_currency_sum_does_not_pick_a_currency(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Equity EURm', 100])
        sheet.append(['Equity USDm', 200])
        result = workbook.create_sheet('Results')
        result.append(['Equity', '=Sheet!B1+Sheet!B2'])
        result.append(['Number of shares (million)', 20])
        result.append(['Target price', '=B1/B2'])
        cells = self.parse(workbook)
        self.assertEqual(cells['Results', 'B1']['unit'], '')
        self.assertEqual(json.loads(cells['Results', 'B1']['metadata_json'])['unit_context']['status'], 'ambiguous')
        self.assertNotIn(cells['Results', 'B3']['unit'], ['EUR/share', 'USD/share'])

    def test_formats_and_unit_declarations_support_currency_and_scale_variants(self):
        workbook = Workbook()
        sheet = workbook.active
        expected = [
            ('Cash in USD billion', 'General', 'USDbn'),
            ('Cash EURbn', 'General', 'EURbn'),
            ('现金（人民币亿元）', 'General', 'CNY_100m'),
            ('Target price', '"HK$"0.00', 'HKD/share'),
            ('Target price', '"US$"0.00', 'USD/share'),
            ('Target price', '[$€-407]0.00', 'EUR/share'),
            ('Cash JPY million', 'General', 'JPYm'),
        ]
        for row, (label, format_code, _) in enumerate(expected, 1):
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, 10).number_format = format_code
        cells = self.parse(workbook)
        for row, (_, _, unit) in enumerate(expected, 1):
            self.assertEqual(cells['Sheet', f'B{row}']['unit'], unit)

    def test_unknown_currency_scale_external_references_and_cycles_stay_unknown(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Target price', 15])
        sheet['B1'].number_format = '$0.00'
        sheet.append(['Number of shares', 911])
        sheet.append(['Target price', "='[unavailable.xlsx]Source'!B3"])
        sheet.append(['Target price', '=B5'])
        sheet.append(['Target price', '=B4'])
        cells = self.parse(workbook)
        self.assertEqual(cells['Sheet', 'B2']['unit'], 'share_count_unspecified_scale')
        for ref in ['B1', 'B3', 'B4', 'B5']:
            self.assertEqual(cells['Sheet', ref]['unit'], 'per_share')

    def test_explicit_source_units_survive_unsupported_formula_functions(self):
        workbook = Workbook()
        source = workbook.active
        source.append(['Revenue EURm', '=SUMIFS(C1:C5,D1:D5,"a")'])
        output = workbook.create_sheet('Output')
        output.append(['Revenue', '=Sheet!B1'])
        cells = self.parse(workbook)
        self.assertEqual(cells['Output', 'B1']['unit'], 'EURm')


if __name__ == '__main__':
    unittest.main()
