"""Small real documents for reader integration tests; no network or model calls."""
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

target = Path(sys.argv[1])
target.mkdir(parents=True, exist_ok=True)
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = 'Valuation'
sheet.append(['Valuation Date', datetime(2026, 8, 31)])
sheet['B1'].number_format = 'yyyy-mm-dd'
for row, label, value in [(3, 'Revenue segment A', 500), (4, 'Revenue segment B', 700),
                          (5, 'Revenue', '=SUM(B3:B4)'), (6, 'Current Price', 100),
                          (7, 'Target Price', '=B5/10'), (8, 'Upside', '=B7/B6-1')]:
    sheet.cell(row, 1, label)
    sheet.cell(row, 2, value)
sheet['B7'].number_format = '"CNY/share" 0.00'
sheet['B8'].number_format = '0.0%'
sheet['A10'] = 'Long original note: ' + 'x' * 5100
sheet['D12'] = 'Window edge'
hidden = workbook.create_sheet('Hidden assumptions')
hidden['A1'] = 'Hidden source'
hidden.sheet_state = 'hidden'
workbook.defined_names.add(DefinedName('Primary_Target_Price', attr_text="'Valuation'!$B$7"))
workbook.save(target / 'model.xlsx')
workbook.close()
ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(target / 'model.xlsx') as original:
    entries = {name: original.read(name) for name in original.namelist()}
xml = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
for ref, value in [('B5', '1200'), ('B7', '120'), ('B8', '0.2')]:
    xml.find(f'.//m:c[@r="{ref}"]/m:v', ns).text = value
entries['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
with zipfile.ZipFile(target / 'model.xlsx', 'w') as updated:
    for name, contents in entries.items():
        updated.writestr(name, contents)

with zipfile.ZipFile(target / 'notes.docx', 'w') as doc:
    doc.writestr('word/document.xml', '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
      <w:p><w:r><w:t>First paragraph</w:t></w:r></w:p><w:p/>
      <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Revenue</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    </w:body></w:document>''')
with zipfile.ZipFile(target / 'slides.pptx', 'w') as slides:
    slides.writestr('ppt/presentation.xml', '''<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>''')
    slides.writestr('ppt/_rels/presentation.xml.rels', '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/></Relationships>''')
    slides.writestr('ppt/slides/slide1.xml', '''<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Slide one revenue</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>''')
