"""Deterministic workbook covering the complete model tool chain."""
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.workbook.defined_name import DefinedName

target = Path(sys.argv[1])
target.parent.mkdir(parents=True, exist_ok=True)
workbook = openpyxl.Workbook()
workbook.properties.created = datetime(2026, 1, 1)
workbook.properties.modified = datetime(2026, 8, 31)
sheet = workbook.active
sheet.title = 'Valuation'
sheet.append(['Valuation Date', datetime(2026, 8, 31)])
sheet['B1'].number_format = 'yyyy-mm-dd'
for row, label, value in [(3, 'Revenue segment A', 500), (4, 'Revenue segment B', 700),
                          (5, 'Revenue', '=SUM(B3:B4)'), (6, 'Current Price', 100),
                          (7, 'Target Price', '=B5/10'), (8, 'Upside', '=B7/B6-1')]:
    sheet.cell(row, 1, label)
    sheet.cell(row, 2, value)
sheet['B7'].number_format = '"CNY/share" 0.00'
sheet['B8'].number_format = '0.0%'
sheet['A10'] = 'Long original note: ' + 'x' * 5100
sheet['D12'] = 'Window edge'
sheet.merge_cells('D3:E3')
sheet['D3'] = 'Merged annotation'
sheet.freeze_panes = 'B2'
sheet.row_dimensions[11].hidden = True
sheet.column_dimensions['F'].hidden = True
hidden = workbook.create_sheet('Hidden assumptions')
hidden['A1'] = 'Hidden source'
hidden['B1'] = 2
hidden['B2'] = '=Valuation!B3'
hidden.sheet_state = 'veryHidden'
workbook.defined_names.add(DefinedName('Primary_Target_Price', attr_text="'Valuation'!$B$7"))
workbook.defined_names.add(DefinedName('Input_Growth', attr_text="'Hidden assumptions'!$B$1"))
edge = workbook.create_sheet('Formula cases')
edge['A1'] = 3
edge['A2'] = '=Input_Growth*A1'
edge['A3'] = '=SUM(Valuation!B3:B4)'
edge['B1'] = ArrayFormula(ref='B1:B2', text='=A1:A2*2')
edge['C1'] = DataTableFormula(ref='C1:D2', dt2D=False, r1='A1')
edge['E1'] = '=A99'
edge['F1'] = '=G1'
edge['G1'] = '=F1'
edge['H1'] = "='[missing.xlsx]Sheet1'!A1"
edge['I1'] = '=SUM(A:A)'
edge['J1'] = '=SUM(Valuation:Formula cases!A1)'
edge['K1'] = '#REF!'
dates = workbook.create_sheet('Date cases')
dates.append(['Valuation date', 'Not confirmed 2026-09-01'])
dates.append(['Market price date', datetime(2026, 8, 30)])
dates.append(['Forecast', '2030E'])
dates.append(['Valuation date', 'not 2026-10-31'])
dates.append(['Valuation date', '2026-08-01 or 2026-08-02'])
dcf = workbook.create_sheet('DCF')
for row in [['WACC', .1], ['Free cash flow', 100], ['Terminal value', 1000], ['DCF value', '=B2/B1+B3'], ['Equity value', '=B4-100'], ['Diluted shares', 10], ['Per share value', '=B5/B6']]:
    dcf.append(row)
dcf['B7'].number_format = '"CNY/share" 0.00'
sotp = workbook.create_sheet('SOTP')
for row in [['Segment A', 500], ['Segment B', 700], ['SOTP value', '=SUM(B1:B2)']]:
    sotp.append(row)
sensitivity = workbook.create_sheet('Sensitivity')
sensitivity.append(['Target Price', '=125'])
sensitivity['B1'].number_format = '"CNY/share" 0.00'
if len(sys.argv) > 2:
    sheet['B3'] = float(sys.argv[2])
workbook.save(target)
workbook.close()
ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(target) as original:
    entries = {name: original.read(name) for name in original.namelist()}
xml = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
for ref, value in [('B5', '1200'), ('B7', '120'), ('B8', '0.2')]:
    xml.find(f'.//m:c[@r="{ref}"]/m:v', ns).text = value
entries['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
with zipfile.ZipFile(target, 'w') as updated:
    for name, contents in entries.items():
        updated.writestr(name, contents)
