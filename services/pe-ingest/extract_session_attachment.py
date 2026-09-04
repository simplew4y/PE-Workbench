#!/usr/bin/env python3
"""Extract a session attachment to a UTF-8 text sidecar."""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pymupdf as fitz
from openpyxl import load_workbook

from pipeline.private_fund_format_adapters import adapt_document


MAX_OUTPUT_CHARS = 500_000


def _pdf_text(path: Path) -> str:
    parts: list[str] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            parts.append(f"## Page {page_number}\n{text or '[No extractable text]'}")
    return "\n\n".join(parts)


def _workbook_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    parts: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows():
                values = ["" if cell.value is None else str(cell.value) for cell in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append("\t".join(values))
            parts.append(f"## Sheet: {worksheet.title}\n" + ("\n".join(rows) or "[Empty sheet]"))
    finally:
        workbook.close()
    return "\n\n".join(parts)


def _adapter_text(path: Path) -> str:
    parts: list[str] = []
    for chunk in adapt_document(path, max_chars=8_000):
        content = str(chunk.get("content") or "").strip()
        source_ref = str(chunk.get("source_ref") or "").strip()
        if content:
            parts.append(f"### {source_ref}\n{content}" if source_ref else content)
    return "\n\n".join(parts)


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(path)
    elif suffix in {".xlsx", ".xlsm"}:
        text = _workbook_text(path)
    else:
        text = _adapter_text(path)
    if len(text) > MAX_OUTPUT_CHARS:
        text = text[:MAX_OUTPUT_CHARS] + "\n\n[Attachment text truncated at 500000 characters]"
    return text


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract_session_attachment.py SOURCE OUTPUT")
    source = Path(sys.argv[1]).resolve(strict=True)
    output = Path(sys.argv[2]).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.{os.getpid()}.tmp")
    temporary.write_text(extract(source), encoding="utf-8")
    os.replace(temporary, output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
